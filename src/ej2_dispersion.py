"""Ejercicio 2: dispersión, asimetría y curtosis con fórmulas de Excel, y análisis de homogeneidad."""
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import xlsxwriter
from scipy import stats

from carga import BASE, cargar_datos
from ej2_graficos import escribir_graficos
from ej2_medidas import escribir_medidas, fmt
from ej2_preguntas import escribir_preguntas
from ej2_tabla import VARIABLE, calcular_tabla, escribir_tabla, recalcular_con_libreoffice
from hoja_datos import escribir_datos, rango

# Convenciones adoptadas (de uso común, no tomadas de las lecturas del curso)
CV_LIMITE = 30          # CV menor que 30 %: homogénea; 30 % o más: heterogénea
ASIM_LIMITE = 0.5       # |asimetría| menor que 0,5: casi simétrica


def calcular(datos):
    """Las cinco medidas con numpy y scipy (varianza y desviación muestrales; asimetría y curtosis como Excel)."""
    media = float(np.mean(datos))
    desv = float(np.std(datos, ddof=1))
    return {
        "Varianza muestral": float(np.var(datos, ddof=1)),
        "Desviación típica muestral": desv,
        "Coeficiente de variación (%)": desv / media * 100,
        "Asimetría": float(stats.skew(datos, bias=False)),
        "Curtosis (exceso)": float(stats.kurtosis(datos, bias=False)),
    }


def f2(x):
    """Número con dos decimales fijos y coma decimal."""
    return f"{x:.2f}".replace(".", ",")


def clasificar(x):
    """Etiquetas según las convenciones adoptadas."""
    cv, a, k = x["Coeficiente de variación (%)"], x["Asimetría"], x["Curtosis (exceso)"]
    homog = "homogénea" if cv < CV_LIMITE else "heterogénea"
    if abs(a) < ASIM_LIMITE:
        forma = "casi simétrica"
    else:
        forma = "asimétrica " + ("positiva" if a > 0 else "negativa")
    sentido = "negativa (cola hacia las velocidades bajas)" if a < 0 else "positiva (cola hacia las velocidades altas)"
    curt = "platicúrtica" if k < -0.0 else ("leptocúrtica" if k > 0 else "mesocúrtica")
    return homog, forma, sentido, curt


def textos(x, media, mediana, tab):
    homog, forma, sentido, curt = clasificar(x)
    cv, a, k = x["Coeficiente de variación (%)"], x["Asimetría"], x["Curtosis (exceso)"]
    assert (homog, forma, curt) == ("heterogénea", "casi simétrica", "platicúrtica"), "El texto supone estas tres etiquetas"
    assert a < 0 and media < mediana
    interp = {
        "Varianza muestral": f"Es el promedio de los cuadrados de las distancias de cada velocidad a la media, y vale {fmt(x['Varianza muestral'])} (km/h)². Como está en unidades al cuadrado, se interpreta mejor con la desviación típica.",
        "Desviación típica muestral": f"Las velocidades se desvían típicamente unos {fmt(x['Desviación típica muestral'])} km/h de la media de {fmt(media)} km/h.",
        "Coeficiente de variación (%)": f"La desviación típica equivale al {fmt(cv)} % de la media. Es la medida que usamos para decidir si la distribución es homogénea.",
        "Asimetría": f"Es {f2(a)}: negativa y de tamaño pequeño. La cola es un poco más larga hacia las velocidades bajas, en línea con que la mediana ({fmt(mediana)}) quede por encima de la media ({fmt(media)}).",
        "Curtosis (exceso)": f"Es {f2(k)}, medida en exceso respecto a la distribución normal (que vale 0). Al ser negativa la distribución es platicúrtica, más aplanada que la normal.",
    }
    n = int(tab.fi.sum())
    analisis = {
        "¿La distribución es homogénea o heterogénea?": (
            f"Es heterogénea. Con la convención adoptada (CV menor que {CV_LIMITE} % es homogénea), el coeficiente de variación de "
            f"{fmt(cv)} % la deja del lado heterogéneo, aunque cerca del límite. Las velocidades cambian de forma apreciable de un accidente a otro, "
            f"y con otro umbral la etiqueta podría cambiar, por eso conviene dar el valor del CV junto con la conclusión."),
        "¿Qué tipo de concentración presentan los datos?": (
            f"La curtosis en exceso de {f2(k)} indica una distribución platicúrtica, es decir, con poca concentración de datos alrededor de la media. "
            f"Las velocidades se extienden a lo largo del rango en vez de agruparse alrededor de un solo pico, algo que también se ve en el histograma: "
            f"ningún intervalo pasa de {fmt(tab.hi_pct.max())} % de los accidentes."),
        "¿Qué tipo de asimetría presentan los datos?": (
            f"La asimetría de {f2(a)} es negativa. Con la convención adoptada (valor absoluto menor que {str(ASIM_LIMITE).replace(".", ",")} es casi simétrica), la distribución es casi simétrica, "
            f"con una cola ligeramente más larga hacia las velocidades bajas. Esto coincide con que la media ({fmt(media)}) sea algo menor que la mediana ({fmt(mediana)})."),
    }
    return interp, analisis


def escribir_dispersion(wb, ws, pos, df, info_medidas, valores=True):
    datos = df[VARIABLE].to_numpy()
    rg = rango(df, VARIABLE)
    tab = calcular_tabla(df[VARIABLE])
    x = calcular(datos)
    media = info_medidas["medidas"]["Media"]
    mediana = info_medidas["medidas"]["Mediana"]
    interp, analisis = textos(x, media, mediana, tab)
    fila_media = info_medidas["celdas"]["Media"][0]
    v = (lambda t: t) if valores else (lambda t: None)

    sub = wb.add_format({"bold": True, "font_size": 13})
    enc = wb.add_format({"bold": True, "bg_color": "#D9E2F3", "border": 1, "align": "center", "valign": "vcenter"})
    med = wb.add_format({"bold": True, "border": 1, "valign": "vcenter", "text_wrap": True})
    res = wb.add_format({"border": 1, "num_format": "0.00", "valign": "vcenter"})
    res4 = wb.add_format({"border": 1, "num_format": "0.0000", "valign": "vcenter"})
    txt = wb.add_format({"border": 1, "text_wrap": True, "valign": "vcenter"})
    chico = wb.add_format({"border": 1, "text_wrap": True, "valign": "vcenter", "font_color": "#595959", "font_size": 9})
    neg = wb.add_format({"bold": True})
    par = wb.add_format({"text_wrap": True, "valign": "top"})

    f = pos["fila_libre_medidas"] + 1
    ws.write(f, 0, "Dispersión, asimetría y curtosis de la velocidad registrada (km/h)", sub)
    f += 2
    ws.write(f, 0, "Medida", enc)
    ws.write(f, 1, "Resultado", enc)
    ws.merge_range(f, 2, f, 8, "Interpretación", enc)
    ws.merge_range(f, 9, f, 12, "Fórmula de Excel", enc)
    fila_std = f + 3   # fila de Excel de la desviación típica (tercera fila de datos)
    filas = [
        ("Varianza muestral", f"=VAR.S({rg})", "VAR.S", res),
        ("Desviación típica muestral", f"=STDEV.S({rg})", "STDEV.S", res),
        ("Coeficiente de variación (%)", f"=B{fila_std}/B{fila_media}*100", "Desviación típica / Media x 100", res),
        ("Asimetría", f"=SKEW({rg})", "SKEW", res4),
        ("Curtosis (exceso)", f"=KURT({rg})", "KURT", res4),
    ]
    celdas = {}
    for i, (nombre, formula, texto_formula, fmt_res) in enumerate(filas, start=1):
        fila = f + i
        ws.set_row(fila, 48 if nombre in ("Asimetría", "Curtosis (exceso)", "Varianza muestral") else 34)
        ws.write(fila, 0, nombre, med)
        ws.write_formula(fila, 1, formula, fmt_res, v(x[nombre]))
        ws.merge_range(fila, 2, fila, 8, interp[nombre], txt)
        ws.merge_range(fila, 9, fila, 12, texto_formula, chico)
        celdas[nombre] = (fila + 1, 2)
    assert celdas["Desviación típica muestral"][0] == fila_std, "La fila de la desviación típica cambió"

    f += len(filas) + 2
    ws.write(f, 0, "Convención adoptada para interpretar", neg)
    ws.merge_range(f + 1, 0, f + 1, 12,
                   f"Usamos la varianza y la desviación típica muestrales (VAR.S y STDEV.S) porque los 400 registros se toman como una muestra de la accidentalidad del Tolima. "
                   f"Para clasificar usamos convenciones de uso común en estadística descriptiva, que no proceden de las lecturas del curso: homogénea si el CV es menor que {CV_LIMITE} %, "
                   f"casi simétrica si el valor absoluto de la asimetría es menor que {str(ASIM_LIMITE).replace(".", ",")}, y platicúrtica si la curtosis en exceso es negativa. SKEW y KURT son las versiones de muestra que calcula Excel.", par)
    ws.set_row(f + 1, 48)
    f += 3
    ws.write(f, 0, "Análisis de la distribución", sub)
    f += 1
    for pregunta, texto in analisis.items():
        ws.write(f, 0, pregunta, neg)
        ws.merge_range(f + 1, 0, f + 1, 12, texto, par)
        ws.set_row(f + 1, 48)
        f += 3
    pos["fila_libre_dispersion"] = f
    return {"celdas": celdas, "valores": x, "interp": interp, "analisis": analisis}


def construir(ruta, df, valores=True):
    # use_future_functions añade el prefijo _xlfn. a VAR.S, STDEV.S, MODE.SNGL, QUARTILE.INC y PERCENTILE.INC
    wb = xlsxwriter.Workbook(str(ruta), {"use_future_functions": True})
    ws = wb.add_worksheet("Ej2 Tabla")
    pos = escribir_tabla(wb, ws, df, valores=valores)
    tab = calcular_tabla(df[VARIABLE])
    escribir_graficos(wb, ws, pos, tab, valores=valores)
    info_q = escribir_preguntas(wb, ws, pos, tab, valores=valores)
    info_m = escribir_medidas(wb, ws, pos, df, valores=valores)
    info = escribir_dispersion(wb, ws, pos, df, info_m, valores=valores)
    info["info_preguntas"], info["info_medidas"] = info_q, info_m
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    escribir_datos(wb, df)
    wb.close()
    return pos, info


def verificar(df):
    """LibreOffice (copia sin valores guardados) contra numpy/scipy y contra pandas."""
    s = df[VARIABLE]
    otro = {
        "Varianza muestral": s.var(), "Desviación típica muestral": s.std(),
        "Coeficiente de variación (%)": s.std() / s.mean() * 100, "Asimetría": s.skew(), "Curtosis (exceso)": s.kurt(),
    }
    with tempfile.TemporaryDirectory() as tmp:
        sin_cache = Path(tmp) / "ej2_dispersion_sin_cache.xlsx"
        pos, info = construir(sin_cache, df, valores=False)
        hoja = openpyxl.load_workbook(recalcular_con_libreoffice(sin_cache, tmp), data_only=True)["Ej2 Tabla"]
        leidos = {k: hoja.cell(*rc).value for k, rc in info["celdas"].items()}
    peor = 0.0
    print(f"{'Medida':<30}{'Excel (LibreOffice)':>20}{'scipy/numpy':>16}{'pandas':>14}{'dif. máx.':>11}")
    for nombre, valor in info["valores"].items():
        d = max(abs(leidos[nombre] - valor), abs(leidos[nombre] - float(otro[nombre])))
        peor = max(peor, d)
        print(f"{nombre:<30}{leidos[nombre]:>20.8f}{valor:>16.8f}{float(otro[nombre]):>14.8f}{d:>11.1e}")
    assert peor < 1e-8, peor
    print(f"Diferencia máxima entre las tres fuentes: {peor:.2e}")


if __name__ == "__main__":
    df = cargar_datos()
    salida = BASE / "output"
    salida.mkdir(exist_ok=True)
    pos, info = construir(salida / "ej2_dispersion_borrador.xlsx", df)
    for n, t in info["interp"].items():
        print(f"{n}: {t}")
    print()
    for n, t in info["analisis"].items():
        print(f"{n}\n{t}\n")
    verificar(df)
