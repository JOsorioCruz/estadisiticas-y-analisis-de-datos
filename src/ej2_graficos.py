"""Ejercicio 2: histograma, diagrama circular y polígono de frecuencia (gráficos nativos de Excel)."""
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import xlsxwriter

from carga import BASE, cargar_datos
from ej2_tabla import AMPLITUD, VARIABLE, _etiqueta, calcular_tabla, escribir_tabla, recalcular_con_libreoffice
from hoja_datos import escribir_datos

COLOR = "#2F5597"
COL_A, COL_D, COL_E = 0, 3, 4          # intervalo, marca de clase y fi en la tabla
COL_AUX = 11                            # columna L: datos auxiliares del polígono


def c(x, d=1):
    """Número con coma decimal, para el texto en español."""
    return f"{x:.{d}f}".replace(".", ",")


def conclusiones(tab):
    """Textos de conclusión de cada gráfico, armados con los valores de la tabla."""
    n = int(tab.fi.sum())
    k = len(tab)
    lab = [_etiqueta(t.Li, t.Ls, i == k - 1) for i, t in enumerate(tab.itertuples())]
    imax, imin = int(tab.fi.idxmax()), int(tab.fi.idxmin())
    assert imin == 2 and imax == 7, "El texto supone valle en el tercer intervalo y pico en el octavo"
    trio = [tab.fi[i:i + 3].sum() for i in range(k - 2)]
    itrio = trio.index(max(trio))
    assert itrio == 5, "El texto supone que el mejor trío de intervalos vecinos es el sexto al octavo"
    centro = (tab.Li.iloc[0] + tab.Ls.iloc[-1]) / 2
    sobre_centro = int(tab.fi[tab.Li >= centro].sum())
    pct = lambda i: c(tab.hi_pct[i])

    histograma = (
        f"El intervalo con la barra más alta es {lab[imax]} km/h, con {tab.fi[imax]} accidentes ({pct(imax)} %), "
        f"y el de la barra más baja es {lab[imin]} km/h, con solo {tab.fi[imin]} ({pct(imin)} %). "
        f"Ese valle es un hueco entre dos intervalos vecinos que tienen {tab.fi[imin-1]} y {tab.fi[imin+1]} accidentes, "
        f"así que la distribución no es simétrica ni tiene un único pico. "
        f"Entre {tab.Li[itrio]:g} y {tab.Ls[itrio+2]:g} km/h se concentran {trio[itrio]} accidentes "
        f"({c(trio[itrio] / n * 100)} %), cuatro de cada diez. "
        f"La base no trae el límite de velocidad de cada vía, por lo que esto muestra dónde se concentran los siniestros y no cuáles velocidades son excesivas."
    )
    circular = (
        f"Cada porción es el porcentaje de accidentes de un intervalo de velocidad. La más grande es {lab[imax]} con "
        f"{pct(imax)} % y la más pequeña es {lab[imin]} con {pct(imin)} %. Como ninguna porción pasa del {pct(imax)} %, "
        f"ninguna franja de velocidad domina sobre las demás. Los tres intervalos vecinos con más accidentes, de "
        f"{tab.Li[itrio]:g} a {tab.Ls[itrio+2]:g} km/h, suman el {c(trio[itrio] / n * 100)} %. "
        f"Con diez porciones es difícil comparar tamaños a simple vista, por eso conviene apoyarse en el histograma."
    )
    poligono = (
        f"El polígono une las marcas de clase y se cierra con dos intervalos vacíos, uno antes y otro después "
        f"(marcas {c(tab.xi.iloc[0] - AMPLITUD)} y {c(tab.xi.iloc[-1] + AMPLITUD)} km/h), por eso empieza y termina en 0. "
        f"Arranca en {tab.fi[0]} accidentes en la marca {c(tab.xi[0])}, cae a {tab.fi[imin]} en la marca {c(tab.xi[imin])} y "
        f"desde ese valle sube, con una pequeña baja en {c(tab.xi[6])}, hasta el pico de {tab.fi[imax]} en {c(tab.xi[imax])} km/h. "
        f"Después baja a {tab.fi[8]} y {tab.fi[9]}. Como el punto medio del rango es {centro:g} km/h y {sobre_centro} accidentes "
        f"({c(sobre_centro / n * 100)} %) se registraron a esa velocidad o más, la distribución se inclina hacia las velocidades altas."
    )
    return {"Conclusión del histograma": histograma, "Conclusión del diagrama circular": circular,
            "Conclusión del polígono de frecuencia": poligono}


def escribir_graficos(wb, ws, pos, tab, valores=True):
    """Agrega el bloque auxiliar del polígono, los tres gráficos y sus conclusiones."""
    hoja = pos["hoja"]
    p, u, h = pos["primera"], pos["ultima"], pos["encabezado"]
    k = u - p + 1
    amp = f"$B${pos['fila_amp']}"
    v = (lambda x: x) if valores else (lambda x: None)

    negrita = wb.add_format({"bold": True})
    enc = wb.add_format({"bold": True, "bg_color": "#D9E2F3", "border": 1, "align": "center", "text_wrap": True})
    f_1 = wb.add_format({"border": 1, "num_format": "0.0"})
    f_int = wb.add_format({"border": 1, "num_format": "0"})

    # Auxiliar del polígono: clase anterior y posterior con fi = 0
    ws.write(h - 1, COL_AUX, "Auxiliar del polígono (clases vacías en los extremos)", negrita)
    ws.write(h, COL_AUX, "Marca xi", enc)
    ws.write(h, COL_AUX + 1, "fi", enc)
    aux0 = h + 1
    ws.write_formula(aux0, COL_AUX, f"=D{p+1}-{amp}", f_1, v(float(tab.xi.iloc[0] - AMPLITUD)))
    ws.write_number(aux0, COL_AUX + 1, 0, f_int)
    for i in range(k):
        ws.write_formula(aux0 + 1 + i, COL_AUX, f"=D{p+1+i}", f_1, v(float(tab.xi[i])))
        ws.write_formula(aux0 + 1 + i, COL_AUX + 1, f"=E{p+1+i}", f_int, v(int(tab.fi[i])))
    ws.write_formula(aux0 + k + 1, COL_AUX, f"=D{u+1}+{amp}", f_1, v(float(tab.xi.iloc[-1] + AMPLITUD)))
    ws.write_number(aux0 + k + 1, COL_AUX + 1, 0, f_int)
    ws.set_column(COL_AUX, COL_AUX + 1, 12)

    # Histograma
    hist = wb.add_chart({"type": "column"})
    hist.add_series({
        "name": "Frecuencia absoluta (fi)",
        "categories": [hoja, p, COL_A, u, COL_A],
        "values": [hoja, p, COL_E, u, COL_E],
        "gap": 0,
        "fill": {"color": COLOR},
        "border": {"color": "#FFFFFF", "width": 1},
        "data_labels": {"value": True},
    })
    hist.set_title({"name": "Histograma de la velocidad registrada en los accidentes", "name_font": {"size": 12}})
    hist.set_x_axis({"name": "Velocidad registrada (km/h)", "num_font": {"rotation": -45}})
    hist.set_y_axis({"name": "Número de accidentes (fi)", "major_gridlines": {"visible": True, "line": {"color": "#D9D9D9"}}})
    hist.set_legend({"none": True})
    hist.set_size({"width": 700, "height": 360})

    # Diagrama circular
    pie = wb.add_chart({"type": "pie"})
    pie.add_series({
        "name": "Porcentaje de accidentes",
        "categories": [hoja, p, COL_A, u, COL_A],
        "values": [hoja, p, COL_E, u, COL_E],
        "data_labels": {"percentage": True, "leader_lines": True, "position": "outside_end", "num_format": "0.0%"},
    })
    pie.set_title({"name": "Distribución porcentual de los accidentes por intervalo de velocidad (km/h)", "name_font": {"size": 12}})
    pie.set_legend({"position": "right"})
    pie.set_size({"width": 700, "height": 400})

    # Polígono de frecuencia (dispersión con líneas rectas sobre las marcas de clase)
    pol = wb.add_chart({"type": "scatter", "subtype": "straight_with_markers"})
    pol.add_series({
        "name": "Frecuencia absoluta (fi)",
        "categories": [hoja, aux0, COL_AUX, aux0 + k + 1, COL_AUX],
        "values": [hoja, aux0, COL_AUX + 1, aux0 + k + 1, COL_AUX + 1],
        "line": {"color": COLOR, "width": 2.25},
        "marker": {"type": "circle", "size": 6, "fill": {"color": COLOR}, "border": {"color": COLOR}},
        "data_labels": {"value": True, "position": "above"},
    })
    pol.set_title({"name": "Polígono de frecuencia de la velocidad registrada", "name_font": {"size": 12}})
    pol.set_x_axis({"name": "Marca de clase: velocidad registrada (km/h)", "min": 20, "max": 150, "major_unit": 10})
    pol.set_y_axis({"name": "Número de accidentes (fi)", "min": 0, "major_gridlines": {"visible": True, "line": {"color": "#D9D9D9"}}})
    pol.set_legend({"none": True})
    pol.set_size({"width": 700, "height": 360})

    # Ubicación: cada gráfico con su conclusión debajo
    tit = wb.add_format({"bold": True})
    txt = wb.add_format({"text_wrap": True, "valign": "top"})
    textos = conclusiones(tab)
    fila = pos["total"] + 4
    pos["inicios_graficos"] = []
    ws.write(fila - 1, 0, "Gráficos estadísticos y conclusiones", wb.add_format({"bold": True, "font_size": 13}))
    for grafico, alto_filas, (titulo, texto) in zip((hist, pie, pol), (19, 21, 19), textos.items()):
        pos["inicios_graficos"].append(fila)
        ws.insert_chart(fila, 0, grafico)
        ws.write(fila + alto_filas, 0, titulo, tit)
        ws.merge_range(fila + alto_filas + 1, 0, fila + alto_filas + 1, 9, texto, txt)
        ws.set_row(fila + alto_filas + 1, 62)
        fila += alto_filas + 3
    pos["fila_libre"] = fila
    return textos


def construir(ruta, df, valores=True):
    wb = xlsxwriter.Workbook(str(ruta))
    ws = wb.add_worksheet("Ej2 Tabla")
    pos = escribir_tabla(wb, ws, df, valores=valores)
    textos = escribir_graficos(wb, ws, pos, calcular_tabla(df[VARIABLE]), valores=valores)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    escribir_datos(wb, df)
    wb.close()
    return pos, textos


def verificar_recalculo(df):
    """Recalcula con LibreOffice una copia sin valores guardados y compara el auxiliar del polígono."""
    tab = calcular_tabla(df[VARIABLE])
    esperado = [(tab.xi.iloc[0] - AMPLITUD, 0)] + list(zip(tab.xi, tab.fi)) + [(tab.xi.iloc[-1] + AMPLITUD, 0)]
    with tempfile.TemporaryDirectory() as tmp:
        sin_cache = Path(tmp) / "ej2_graficos_sin_cache.xlsx"
        pos, _ = construir(sin_cache, df, valores=False)
        hoja = openpyxl.load_workbook(recalcular_con_libreoffice(sin_cache, tmp), data_only=True)["Ej2 Tabla"]
        fila0 = pos["encabezado"] + 2
        dif = max(abs(hoja.cell(fila0 + i, COL_AUX + 1).value - x) + abs(hoja.cell(fila0 + i, COL_AUX + 2).value - y)
                  for i, (x, y) in enumerate(esperado))
    print(f"Auxiliar del polígono (12 puntos), LibreOffice vs numpy: diferencia máxima = {dif:.3e}")


def verificar(ruta):
    """Comprueba que los gráficos son nativos y que apuntan a las celdas correctas."""
    with zipfile.ZipFile(ruta) as z:
        nombres = z.namelist()
        graficos = sorted(n for n in nombres if n.startswith("xl/charts/chart"))
        imagenes = [n for n in nombres if n.startswith("xl/media/")]
        print("Gráficos nativos:", graficos)
        print("Imágenes incrustadas:", imagenes or "ninguna")
        for g in graficos:
            xml = z.read(g).decode("utf-8")
            tipo = next(t for t in ("barChart", "pieChart", "scatterChart") if t in xml)
            refs = sorted(set(__import__("re").findall(r"<c:f>([^<]+)</c:f>", xml)))
            print(f"  {g}: {tipo}; referencias: {refs}")


if __name__ == "__main__":
    df = cargar_datos()
    salida = BASE / "output"
    salida.mkdir(exist_ok=True)
    ruta = salida / "ej2_graficos_borrador.xlsx"
    pos, textos = construir(ruta, df)
    for t, x in textos.items():
        print(f"\n{t}:\n{x}")
    print()
    verificar(ruta)
    verificar_recalculo(df)
