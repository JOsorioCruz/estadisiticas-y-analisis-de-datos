"""Ejercicio 2: medidas de tendencia central y de posición con fórmulas de Excel, e interpretación."""
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import xlsxwriter

from carga import BASE, cargar_datos
from ej2_graficos import escribir_graficos
from ej2_preguntas import escribir_preguntas
from ej2_tabla import VARIABLE, calcular_tabla, escribir_tabla, recalcular_con_libreoffice
from hoja_datos import escribir_datos, rango


def fmt(x):
    """Número con coma decimal y hasta dos decimales, sin ceros sobrantes."""
    return f"{round(float(x), 2):g}".replace(".", ",")


def moda_excel(datos):
    """MODE.SNGL devuelve el valor más frecuente y, si hay empate, el que aparece primero en los datos."""
    conteo = pd.Series(datos).value_counts()
    empatados = set(conteo[conteo == conteo.max()].index)
    return next(x for x in datos if x in empatados), int(conteo.max())


# (nombre, texto de la fórmula, función que genera la fórmula de Excel dado el rango, cálculo en Python)
def definir_medidas(datos):
    s = pd.Series(datos)
    moda, veces = moda_excel(list(datos))
    return [
        ("Moda", "MODE.SNGL", lambda rg: f"=MODE.SNGL({rg})", moda),
        ("Media", "AVERAGE", lambda rg: f"=AVERAGE({rg})", float(s.mean())),
        ("Mediana", "MEDIAN", lambda rg: f"=MEDIAN({rg})", float(s.median())),
        ("Mínimo", "MIN", lambda rg: f"=MIN({rg})", float(s.min())),
        ("Máximo", "MAX", lambda rg: f"=MAX({rg})", float(s.max())),
        ("Rango", "Máximo - Mínimo", None, float(s.max() - s.min())),
        ("Cuartil 1 (Q1)", "QUARTILE.INC(rango, 1)", lambda rg: f"=QUARTILE.INC({rg},1)", float(np.percentile(datos, 25))),
        ("Decil 5 (D5)", "PERCENTILE.INC(rango, 0.5)", lambda rg: f"=PERCENTILE.INC({rg},0.5)", float(np.percentile(datos, 50))),
        ("Cuartil 3 (Q3)", "QUARTILE.INC(rango, 3)", lambda rg: f"=QUARTILE.INC({rg},3)", float(np.percentile(datos, 75))),
        ("Percentil 10 (P10)", "PERCENTILE.INC(rango, 0.1)", lambda rg: f"=PERCENTILE.INC({rg},0.1)", float(np.percentile(datos, 10))),
        ("Percentil 80 (P80)", "PERCENTILE.INC(rango, 0.8)", lambda rg: f"=PERCENTILE.INC({rg},0.8)", float(np.percentile(datos, 80))),
    ], veces


def interpretaciones(m, veces, intervalo_modal, n):
    """Frase de interpretación por medida, con los valores calculados."""
    x = {nombre: valor for nombre, _, _, valor in m}
    mo, me, md, mn, mx = x["Moda"], x["Media"], x["Mediana"], x["Mínimo"], x["Máximo"]
    return {
        "Moda": f"La velocidad que más se repite en los accidentes es {fmt(mo)} km/h, con {veces} registros. Al ser una variable continua este valor pesa poco: el intervalo más frecuente de la tabla es {intervalo_modal} km/h.",
        "Media": f"En promedio, el vehículo causante del accidente iba a {fmt(me)} km/h.",
        "Mediana": f"La mitad de los accidentes ocurrió a {fmt(md)} km/h o menos y la otra mitad a más de {fmt(md)} km/h. Queda por encima de la media ({fmt(me)}), lo que apunta a una ligera cola hacia las velocidades bajas.",
        "Mínimo": f"La velocidad más baja registrada en un accidente fue de {fmt(mn)} km/h.",
        "Máximo": f"La velocidad más alta registrada en un accidente fue de {fmt(mx)} km/h.",
        "Rango": f"Entre el accidente de menor velocidad y el de mayor velocidad hay {fmt(mx - mn)} km/h de diferencia.",
        "Cuartil 1 (Q1)": f"El 25 % de los accidentes ocurrió a velocidades entre {fmt(mn)} y {fmt(x['Cuartil 1 (Q1)'])} km/h.",
        "Decil 5 (D5)": f"El decil 5 deja el 50 % de los accidentes a {fmt(x['Decil 5 (D5)'])} km/h o menos, por eso coincide con la mediana.",
        "Cuartil 3 (Q3)": f"El 75 % de los accidentes ocurrió a velocidades entre {fmt(mn)} y {fmt(x['Cuartil 3 (Q3)'])} km/h, y solo el 25 % superó los {fmt(x['Cuartil 3 (Q3)'])} km/h.",
        "Percentil 10 (P10)": f"El 10 % de los accidentes ocurrió a velocidades entre {fmt(mn)} y {fmt(x['Percentil 10 (P10)'])} km/h.",
        "Percentil 80 (P80)": f"El 80 % de los accidentes ocurrió a {fmt(x['Percentil 80 (P80)'])} km/h o menos, y el 20 % restante superó esa velocidad.",
    }


def escribir_medidas(wb, ws, pos, df, valores=True):
    datos = df[VARIABLE].to_numpy()
    rg = rango(df, VARIABLE)
    tab = calcular_tabla(df[VARIABLE])
    m, veces = definir_medidas(datos)
    intervalo_modal = f"[{tab.Li[tab.fi.idxmax()]:g} - {tab.Ls[tab.fi.idxmax()]:g})"
    textos = interpretaciones(m, veces, intervalo_modal, len(datos))
    v = (lambda x: x) if valores else (lambda x: None)

    sub = wb.add_format({"bold": True, "font_size": 13})
    enc = wb.add_format({"bold": True, "bg_color": "#D9E2F3", "border": 1, "align": "center", "valign": "vcenter"})
    med = wb.add_format({"bold": True, "border": 1, "valign": "vcenter"})
    res = wb.add_format({"border": 1, "num_format": "0.00", "valign": "vcenter"})
    txt = wb.add_format({"border": 1, "text_wrap": True, "valign": "vcenter"})
    chico = wb.add_format({"border": 1, "text_wrap": True, "valign": "vcenter", "font_color": "#595959", "font_size": 9})

    f = pos["fila_libre_preguntas"]
    ws.write(f, 0, "Medidas de tendencia central y de posición de la velocidad registrada (km/h)", sub)
    f += 2
    ws.write(f, 0, "Medida", enc)
    ws.write(f, 1, "Resultado", enc)
    ws.merge_range(f, 2, f, 8, "Interpretación", enc)
    ws.merge_range(f, 9, f, 12, "Fórmula de Excel", enc)
    celdas = {}
    fila_de = {}
    for i, (nombre, texto_formula, generador, valor) in enumerate(m, start=1):
        fila = f + i
        fila_de[nombre] = fila + 1
        ws.set_row(fila, 34)
        ws.write(fila, 0, nombre, med)
        if generador is None:  # Rango = máximo - mínimo, apuntando a las celdas de la propia tabla
            formula = f"=B{fila_de['Máximo']}-B{fila_de['Mínimo']}"
        else:
            formula = generador(rg)
        ws.write_formula(fila, 1, formula, res, v(valor))
        ws.merge_range(fila, 2, fila, 8, textos[nombre], txt)
        ws.merge_range(fila, 9, fila, 12, texto_formula, chico)
        celdas[nombre] = (fila + 1, 2)
    pos["fila_libre_medidas"] = f + len(m) + 2
    return {"celdas": celdas, "textos": textos, "medidas": {n: val for n, _, _, val in m}, "veces_moda": veces}


def construir(ruta, df, valores=True):
    # use_future_functions añade el prefijo _xlfn. a MODE.SNGL, QUARTILE.INC y PERCENTILE.INC; sin él Excel muestra #NAME?
    wb = xlsxwriter.Workbook(str(ruta), {"use_future_functions": True})
    ws = wb.add_worksheet("Ej2 Tabla")
    pos = escribir_tabla(wb, ws, df, valores=valores)
    tab = calcular_tabla(df[VARIABLE])
    escribir_graficos(wb, ws, pos, tab, valores=valores)
    escribir_preguntas(wb, ws, pos, tab, valores=valores)
    info = escribir_medidas(wb, ws, pos, df, valores=valores)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    escribir_datos(wb, df)
    wb.close()
    return pos, info


def verificar(df):
    """Recalcula con LibreOffice una copia sin valores guardados y compara con numpy y con pandas."""
    s = df[VARIABLE]
    otro = {  # segunda implementación independiente (pandas) de las mismas medidas
        "Moda": moda_excel(list(s))[0], "Media": s.mean(), "Mediana": s.median(), "Mínimo": s.min(), "Máximo": s.max(),
        "Rango": s.max() - s.min(), "Cuartil 1 (Q1)": s.quantile(.25), "Decil 5 (D5)": s.quantile(.5),
        "Cuartil 3 (Q3)": s.quantile(.75), "Percentil 10 (P10)": s.quantile(.1), "Percentil 80 (P80)": s.quantile(.8),
    }
    with tempfile.TemporaryDirectory() as tmp:
        sin_cache = Path(tmp) / "ej2_medidas_sin_cache.xlsx"
        pos, info = construir(sin_cache, df, valores=False)
        hoja = openpyxl.load_workbook(recalcular_con_libreoffice(sin_cache, tmp), data_only=True)["Ej2 Tabla"]
        leidos = {k: hoja.cell(*rc).value for k, rc in info["celdas"].items()}
    peor = 0.0
    print(f"{'Medida':<22}{'Excel (LibreOffice)':>22}{'numpy':>14}{'pandas':>14}{'dif. máx.':>12}")
    for nombre, valor in info["medidas"].items():
        d = max(abs(leidos[nombre] - valor), abs(leidos[nombre] - float(otro[nombre])))
        peor = max(peor, d)
        print(f"{nombre:<22}{leidos[nombre]:>22.6f}{valor:>14.6f}{float(otro[nombre]):>14.6f}{d:>12.1e}")
    assert peor < 1e-9, peor
    print(f"Diferencia máxima entre las tres fuentes: {peor:.2e}")


if __name__ == "__main__":
    df = cargar_datos()
    salida = BASE / "output"
    salida.mkdir(exist_ok=True)
    pos, info = construir(salida / "ej2_medidas_borrador.xlsx", df)
    for n, t in info["textos"].items():
        print(f"{n}: {t}")
    print()
    verificar(df)
