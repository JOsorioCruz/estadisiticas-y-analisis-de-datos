"""Ejercicio 2: intervalo de mayor frecuencia e intervalo donde se acumula el 60 %."""
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import xlsxwriter

from carga import BASE, cargar_datos
from ej2_graficos import c, escribir_graficos
from ej2_tabla import AMPLITUD, VARIABLE, _etiqueta, calcular_tabla, escribir_tabla, recalcular_con_libreoffice
from hoja_datos import escribir_datos

OBJETIVO = 0.60   # porcentaje acumulado que pide el enunciado


def respuestas(tab, objetivo=OBJETIVO):
    """Respuestas calculadas con numpy, independientes de las fórmulas de Excel."""
    k = len(tab)
    lab = [_etiqueta(t.Li, t.Ls, i == k - 1) for i, t in enumerate(tab.itertuples())]
    imax = int(np.argmax(tab.fi.to_numpy()))
    m = int((tab.Hi.to_numpy() < objetivo).sum())    # intervalos con acumulado por debajo del objetivo
    hi_ant = float(tab.Hi[m - 1])
    hi_act = float(tab.Hi[m])
    estimacion = float(tab.Li[m] + (objetivo - hi_ant) / (hi_act - hi_ant) * AMPLITUD)
    return {"lab": lab, "imax": imax, "m": m, "hi_ant": hi_ant, "hi_act": hi_act, "estimacion": estimacion}


def texto_respuestas(tab, r):
    n = int(tab.fi.sum())
    i, m = r["imax"], r["m"]
    lab = r["lab"]
    p1 = (f"El intervalo con más accidentes es {lab[i]} km/h, con {tab.fi[i]} de los {n} registros ({c(tab.hi_pct[i])} %). "
          f"Es el intervalo modal de la tabla.")
    p2 = (f"El 60 % se acumula en el intervalo {lab[m]} km/h. Hasta {lab[m-1]} el acumulado llega a {c(tab.Hi_pct[m-1])} % "
          f"({tab.Fi[m-1]} accidentes), todavía por debajo del 60 %, y al sumar los {tab.fi[m]} accidentes de {lab[m]} pasa a "
          f"{c(tab.Hi_pct[m])} % ({tab.Fi[m]} accidentes). Es decir que el {c(tab.Hi_pct[m], 0)} % de los accidentes ocurrió a menos de "
          f"{tab.Ls[m]:g} km/h. Si suponemos que los datos se reparten de forma pareja dentro del intervalo, el 60 % se alcanza cerca de "
          f"los {c(r['estimacion'])} km/h.")
    return p1, p2


def escribir_preguntas(wb, ws, pos, tab, valores=True):
    hoja_p, u = pos["primera"] + 1, pos["ultima"] + 1          # filas de Excel del cuerpo de la tabla
    A, B, E, H, I, J = (f"${col}${hoja_p}:${col}${u}" for col in "ABEHIJ")
    r = respuestas(tab)
    v = (lambda x: x) if valores else (lambda x: None)
    n = int(tab.fi.sum())
    p1, p2 = texto_respuestas(tab, r)

    negrita = wb.add_format({"bold": True})
    sub = wb.add_format({"bold": True, "font_size": 13})
    ent = wb.add_format({"bg_color": "#FFF2CC", "border": 1, "num_format": "0%"})
    f_txt = wb.add_format({"border": 1, "align": "right"})
    f_int = wb.add_format({"border": 1, "num_format": "0"})
    f_2 = wb.add_format({"border": 1, "num_format": "0.00"})
    f_1 = wb.add_format({"border": 1, "num_format": "0.0"})
    txt = wb.add_format({"text_wrap": True, "valign": "top"})

    f = pos["fila_libre"] + 1
    ws.write(f, 0, "Preguntas con base en la tabla de frecuencia", sub)
    f += 2
    ws.write(f, 0, "1. ¿En cuál intervalo se observa la mayor cantidad de datos?", negrita)
    celdas = {}
    filas = [
        ("Intervalo de mayor frecuencia", f"=INDEX({A},MATCH(MAX({E}),{E},0))", r["lab"][r["imax"]], f_txt),
        ("Frecuencia absoluta máxima (fi)", f"=MAX({E})", int(tab.fi.max()), f_int),
        ("Porcentaje de los accidentes (hi %)", f"=INDEX({I},MATCH(MAX({E}),{E},0))", float(tab.hi_pct[r["imax"]]), f_2),
    ]
    for j, (etq, formula, val, fmt) in enumerate(filas, start=1):
        ws.write(f + j, 0, etq)
        ws.write_formula(f + j, 1, formula, fmt, v(val))
        celdas[etq] = (f + j + 1, 2)      # fila y columna de Excel (base 1)
    ws.merge_range(f + 4, 0, f + 4, 9, p1, txt)
    ws.set_row(f + 4, 34)

    f += 6
    ws.write(f, 0, "2. ¿Hasta cuál intervalo se acumula el 60 % de los datos?", negrita)
    ws.write(f + 1, 0, "Porcentaje acumulado buscado")
    ws.write_number(f + 1, 1, OBJETIVO, ent)
    obj = f"$B${f + 2}"
    m = f"COUNTIF({H},\"<\"&{obj})"
    filas = [
        ("Intervalo donde se alcanza el 60 %", f"=INDEX({A},{m}+1)", r["lab"][r["m"]], f_txt),
        ("Hi (%) del intervalo anterior", f"=INDEX({J},{m})", float(tab.Hi_pct[r["m"] - 1]), f_2),
        ("Hi (%) del intervalo donde se alcanza", f"=INDEX({J},{m}+1)", float(tab.Hi_pct[r["m"]]), f_2),
        ("Estimación dentro del intervalo (km/h)",
         f"=INDEX($B${hoja_p}:$B${u},{m}+1)+({obj}-INDEX({H},{m}))/(INDEX({H},{m}+1)-INDEX({H},{m}))*$B${pos['fila_amp']}",
         r["estimacion"], f_1),
    ]
    for j, (etq, formula, val, fmt) in enumerate(filas, start=2):
        ws.write(f + j, 0, etq)
        ws.write_formula(f + j, 1, formula, fmt, v(val))
        celdas[etq] = (f + j + 1, 2)
    ws.merge_range(f + 6, 0, f + 6, 9, p2, txt)
    ws.set_row(f + 6, 62)
    pos["fila_libre_preguntas"] = f + 8
    return {"celdas": celdas, "textos": (p1, p2)}


def construir(ruta, df, valores=True):
    wb = xlsxwriter.Workbook(str(ruta))
    ws = wb.add_worksheet("Ej2 Tabla")
    pos = escribir_tabla(wb, ws, df, valores=valores)
    tab = calcular_tabla(df[VARIABLE])
    escribir_graficos(wb, ws, pos, tab, valores=valores)
    info = escribir_preguntas(wb, ws, pos, tab, valores=valores)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    escribir_datos(wb, df)
    wb.close()
    return pos, info


def verificar(df):
    """Recalcula con LibreOffice una copia sin valores guardados y compara con numpy."""
    datos = df[VARIABLE].to_numpy()
    tab = calcular_tabla(df[VARIABLE])
    r = respuestas(tab)
    # Comprobación independiente desde los datos crudos, sin pasar por la tabla
    conteos, _ = np.histogram(datos, bins=np.arange(32, 143, 11))
    acum = np.cumsum(conteos) / len(datos)
    assert int(np.argmax(conteos)) == r["imax"] and int((acum < OBJETIVO).sum()) == r["m"]
    print(f"Datos crudos: fi máxima = {conteos.max()} en el intervalo {r['lab'][r['imax']]}; "
          f"acumulado antes = {acum[r['m']-1]:.4f}, después = {acum[r['m']]:.4f}")
    print(f"Percentil 60 exacto de los datos crudos (numpy, referencia): {np.percentile(datos, 60):.2f} km/h")
    with tempfile.TemporaryDirectory() as tmp:
        sin_cache = Path(tmp) / "ej2_preguntas_sin_cache.xlsx"
        pos, info = construir(sin_cache, df, valores=False)
        hoja = openpyxl.load_workbook(recalcular_con_libreoffice(sin_cache, tmp), data_only=True)["Ej2 Tabla"]
        leidos = {k: hoja.cell(*rc).value for k, rc in info["celdas"].items()}
    esperados = {
        "Intervalo de mayor frecuencia": r["lab"][r["imax"]],
        "Frecuencia absoluta máxima (fi)": int(tab.fi.max()),
        "Porcentaje de los accidentes (hi %)": float(tab.hi_pct[r["imax"]]),
        "Intervalo donde se alcanza el 60 %": r["lab"][r["m"]],
        "Hi (%) del intervalo anterior": float(tab.Hi_pct[r["m"] - 1]),
        "Hi (%) del intervalo donde se alcanza": float(tab.Hi_pct[r["m"]]),
        "Estimación dentro del intervalo (km/h)": r["estimacion"],
    }
    for k, esp in esperados.items():
        leido = leidos[k]
        ok = leido == esp if isinstance(esp, str) else abs(leido - esp) < 1e-9
        print(f"  {'OK ' if ok else 'ERROR'} {k}: Excel(LibreOffice) = {leido}; Python = {esp}")
        assert ok, k


if __name__ == "__main__":
    df = cargar_datos()
    salida = BASE / "output"
    salida.mkdir(exist_ok=True)
    pos, info = construir(salida / "ej2_preguntas_borrador.xlsx", df)
    for t in info["textos"]:
        print(t + "\n")
    verificar(df)
