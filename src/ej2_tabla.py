"""Ejercicio 2: tabla de frecuencia agrupada de la velocidad registrada (km/h)."""
import math
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import xlsxwriter
import xlsxwriter.utility as xu

from carga import BASE, RUTA_ANEXO, cargar_datos
from hoja_datos import HOJA, escribir_datos, rango

VARIABLE = "VELOCIDAD REGISTRADA (km/H)"
LIM_INICIAL = 32   # límite inferior del primer intervalo, un número redondo por debajo del mínimo (32,1)
AMPLITUD = 11      # A = R/k = 108,7/10 = 10,87, redondeada hacia arriba para cubrir el máximo (140,8)
SOFFICE = "/opt/homebrew/bin/soffice"

ENCABEZADOS = ["Intervalo", "Límite inferior (Li)", "Límite superior (Ls)", "Marca de clase (xi)",
               "fi", "Fi", "hi", "Hi", "hi (%)", "Hi (%)"]


def numero_intervalos(n):
    """Regla de Sturges: k = 1 + log2(n), redondeada hacia arriba."""
    return math.ceil(1 + math.log2(n))


def calcular_tabla(datos, lim_inicial=LIM_INICIAL, amplitud=AMPLITUD):
    """Misma tabla que las fórmulas de Excel, calculada con numpy."""
    n = len(datos)
    k = numero_intervalos(n)
    li = np.array([lim_inicial + i * amplitud for i in range(k)], dtype=float)
    ls = li + amplitud
    fi = np.array([
        int(((datos >= a) & ((datos <= b) if i == k - 1 else (datos < b))).sum())
        for i, (a, b) in enumerate(zip(li, ls))
    ])
    tabla = pd.DataFrame({"Li": li, "Ls": ls, "xi": (li + ls) / 2, "fi": fi})
    tabla["Fi"] = tabla["fi"].cumsum()
    tabla["hi"] = tabla["fi"] / n
    tabla["Hi"] = tabla["Fi"] / n
    tabla["hi_pct"] = tabla["hi"] * 100
    tabla["Hi_pct"] = tabla["Hi"] * 100
    return tabla


def _etiqueta(a, b, ultimo):
    return f"[{a:g} - {b:g}{']' if ultimo else ')'}"


def escribir_tabla(wb, ws, df, valores=True, fila0=0):
    """Escribe parámetros y tabla con fórmulas. Con valores=False no se guarda el valor calculado
    (se usa para comprobar que LibreOffice recalcula las fórmulas por su cuenta).
    Devuelve las posiciones de la tabla para que las use la tarea de gráficos."""
    datos = df[VARIABLE]
    rg = rango(df, VARIABLE)
    n = len(datos)
    k = numero_intervalos(n)
    tab = calcular_tabla(datos)
    v = (lambda x: x) if valores else (lambda x: None)

    negrita = wb.add_format({"bold": True})
    titulo = wb.add_format({"bold": True, "font_size": 13})
    enc = wb.add_format({"bold": True, "bg_color": "#D9E2F3", "border": 1, "align": "center",
                         "valign": "vcenter", "text_wrap": True})
    ent = wb.add_format({"bg_color": "#FFF2CC", "border": 1})
    f_txt = wb.add_format({"border": 1, "align": "center"})
    f_int = wb.add_format({"border": 1, "num_format": "0"})
    f_1 = wb.add_format({"border": 1, "num_format": "0.0"})
    f_4 = wb.add_format({"border": 1, "num_format": "0.0000"})
    f_2 = wb.add_format({"border": 1, "num_format": "0.00"})
    tot = wb.add_format({"bold": True, "top": 2, "border": 1})
    tot_int = wb.add_format({"bold": True, "border": 1, "num_format": "0"})
    tot_4 = wb.add_format({"bold": True, "border": 1, "num_format": "0.0000"})
    tot_2 = wb.add_format({"bold": True, "border": 1, "num_format": "0.00"})

    ws.write(fila0, 0, "Tabla de frecuencia con datos agrupados: velocidad registrada (km/h)", titulo)

    # Parámetros de construcción (filas de Excel, base 1)
    r = fila0 + 3
    par = [
        ("Número de datos (n)", f"=COUNT({rg})", n, f_int),
        ("Valor mínimo", f"=MIN({rg})", float(datos.min()), f_1),
        ("Valor máximo", f"=MAX({rg})", float(datos.max()), f_1),
        ("Rango (R = máximo - mínimo)", f"=B{r+2}-B{r+1}", float(datos.max() - datos.min()), f_1),
        ("Número de intervalos (k, regla de Sturges)", f"=ROUNDUP(1+LOG(B{r},2),0)", k, f_int),
        ("Amplitud teórica (R / k)", f"=B{r+3}/B{r+4}", float((datos.max() - datos.min()) / k), f_2),
    ]
    for i, (etq, formula, valor, fmt) in enumerate(par):
        ws.write(fila0 + 2 + i, 0, etq, negrita)
        ws.write_formula(fila0 + 2 + i, 1, formula, fmt, v(valor))
    ws.write(fila0 + 8, 0, "Amplitud adoptada (A), redondeada hacia arriba", negrita)
    ws.write_number(fila0 + 8, 1, AMPLITUD, ent)
    ws.write(fila0 + 9, 0, "Límite inferior del primer intervalo", negrita)
    ws.write_number(fila0 + 9, 1, LIM_INICIAL, ent)
    fila_amp, fila_ini = fila0 + 9, fila0 + 10   # filas de Excel de A y del límite inicial

    # Encabezado y cuerpo de la tabla
    h = fila0 + 11
    ws.set_row(h, 32)
    for j, e in enumerate(ENCABEZADOS):
        ws.write(h, j, e, enc)
    p, u = h + 1, h + k   # primera y última fila (base 0) del cuerpo
    for i in range(k):
        f = p + i
        x = f + 1  # fila de Excel
        ultimo = i == k - 1
        t = tab.iloc[i]
        ws.write(f, 0, _etiqueta(t.Li, t.Ls, ultimo), f_txt)
        ws.write_formula(f, 1, f"=B{fila_ini}" if i == 0 else f"=C{x-1}", f_1, v(t.Li))
        ws.write_formula(f, 2, f"=B{x}+$B${fila_amp}", f_1, v(t.Ls))
        ws.write_formula(f, 3, f"=(B{x}+C{x})/2", f_1, v(t.xi))
        op = '"<="' if ultimo else '"<"'
        ws.write_formula(f, 4, f'=COUNTIFS({rg},">="&B{x},{rg},{op}&C{x})', f_int, v(int(t.fi)))
        ws.write_formula(f, 5, f"=E{x}" if i == 0 else f"=F{x-1}+E{x}", f_int, v(int(t.Fi)))
        ws.write_formula(f, 6, f"=E{x}/$B${fila0+3}", f_4, v(t.hi))
        ws.write_formula(f, 7, f"=F{x}/$B${fila0+3}", f_4, v(t.Hi))
        ws.write_formula(f, 8, f"=G{x}*100", f_2, v(t.hi_pct))
        ws.write_formula(f, 9, f"=H{x}*100", f_2, v(t.Hi_pct))
    t_ = u + 1
    a, b = p + 1, u + 1  # filas de Excel del cuerpo
    ws.write(t_, 0, "Total", tot)
    for j in (1, 2, 3):
        ws.write_blank(t_, j, None, tot)
    ws.write_formula(t_, 4, f"=SUM(E{a}:E{b})", tot_int, v(int(tab.fi.sum())))
    ws.write_blank(t_, 5, None, tot)
    ws.write_formula(t_, 6, f"=SUM(G{a}:G{b})", tot_4, v(float(tab.hi.sum())))
    ws.write_blank(t_, 7, None, tot)
    ws.write_formula(t_, 8, f"=SUM(I{a}:I{b})", tot_2, v(float(tab.hi_pct.sum())))
    ws.write_blank(t_, 9, None, tot)
    ws.set_column(0, 0, 44)
    ws.set_column(1, 9, 14)
    ws.write(t_ + 2, 0, "Los intervalos son cerrados a la izquierda y abiertos a la derecha, salvo el último, que incluye su límite superior.")
    return {"primera": p, "ultima": u, "total": t_, "encabezado": h, "fila_amp": fila_amp, "hoja": ws.get_name()}


def construir(ruta, df, valores=True):
    wb = xlsxwriter.Workbook(str(ruta))
    ws = wb.add_worksheet("Ej2 Tabla")
    pos = escribir_tabla(wb, ws, df, valores=valores)
    escribir_datos(wb, df)
    wb.close()
    return pos


XCU_RECALCULO = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def recalcular_con_libreoffice(ruta, carpeta_salida):
    """Convierte a xlsx con un perfil temporal que obliga a recalcular todas las fórmulas al abrir."""
    with tempfile.TemporaryDirectory() as perfil:
        usuario = Path(perfil) / "user"
        usuario.mkdir()
        (usuario / "registrymodifications.xcu").write_text(XCU_RECALCULO, encoding="utf-8")
        subprocess.run([SOFFICE, f"-env:UserInstallation=file://{perfil}", "--headless",
                        "--convert-to", "xlsx", "--outdir", str(carpeta_salida), str(ruta)],
                       check=True, capture_output=True)
    return Path(carpeta_salida) / Path(ruta).name


def verificar(df, pos):
    """Compara: (1) tabla en Excel recalculada por LibreOffice contra numpy, (2) hoja DATOS contra el Anexo 1."""
    tab = calcular_tabla(df[VARIABLE])
    with tempfile.TemporaryDirectory() as tmp:
        sin_cache = Path(tmp) / "ej2_tabla_sin_cache.xlsx"
        construir(sin_cache, df, valores=False)
        recalculado = recalcular_con_libreoffice(sin_cache, tmp)
        hoja = openpyxl.load_workbook(recalculado, data_only=True)["Ej2 Tabla"]
        cols = {2: "Li", 3: "Ls", 4: "xi", 5: "fi", 6: "Fi", 7: "hi", 8: "Hi", 9: "hi_pct", 10: "Hi_pct"}
        dif_max = 0.0
        for i in range(len(tab)):
            for c, nombre in cols.items():
                dif_max = max(dif_max, abs(hoja.cell(pos["primera"] + 1 + i, c).value - tab.iloc[i][nombre]))
        total_fi = hoja.cell(pos["total"] + 1, 5).value
        datos = df[VARIABLE]
        esperados = [len(datos), datos.min(), datos.max(), datos.max() - datos.min(), numero_intervalos(len(datos)),
                     (datos.max() - datos.min()) / numero_intervalos(len(datos))]
        params = [hoja.cell(3 + i, 2).value for i in range(6)]
        dif_max = max(dif_max, *(abs(a - b) for a, b in zip(params, esperados)))
    print(f"LibreOffice vs numpy: diferencia máxima = {dif_max:.3e}; suma de fi = {total_fi}")
    print("Parámetros en Excel (n, mín, máx, R, k, R/k):", [round(x, 4) for x in params])

    # Hoja DATOS del libro contra el Anexo 1, celda por celda
    ws_anexo = openpyxl.load_workbook(RUTA_ANEXO)["DATOS"]
    ws_libro = openpyxl.load_workbook(BASE / "output" / "ej2_tabla_borrador.xlsx")[HOJA]
    filas_a = list(ws_anexo.iter_rows(values_only=True))
    filas_l = list(ws_libro.iter_rows(values_only=True))
    print("Hoja DATOS idéntica al Anexo 1:", filas_a == filas_l, f"({len(filas_l)} filas incluyendo encabezado)")
    return dif_max


if __name__ == "__main__":
    df = cargar_datos()
    salida = BASE / "output"
    salida.mkdir(exist_ok=True)
    pos = construir(salida / "ej2_tabla_borrador.xlsx", df)
    print(calcular_tabla(df[VARIABLE]).round(4).to_string())
    verificar(df, pos)
