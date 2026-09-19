"""Ejercicio 3: coeficiente de determinación, confiabilidad, r de Pearson y nivel de correlación."""
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import xlsxwriter
from scipy import stats

from carga import BASE, cargar_datos
from ej2_tabla import recalcular_con_libreoffice
from ej3_regresion import HOJA, X, Y, escribir_regresion, fmt
from hoja_datos import escribir_datos, rango

# Escala adoptada (convención de uso común, no tomada de las lecturas del curso): (límite inferior de |r|, nivel)
ESCALA = [(1.0, "Perfecta"), (0.9, "Excelente"), (0.8, "Aceptable"), (0.5, "Regular"), (0.0, "Mínima")]


def nivel(r):
    a = abs(r)
    if a == 0:
        return "Sin correlación"
    if a == 1:
        return "Perfecta"
    return next(n for lim, n in ESCALA[1:] if a >= lim)


def calcular(df):
    x, y = df[X].to_numpy(), df[Y].to_numpy()
    r = float(np.corrcoef(x, y)[0, 1])
    return {"r": r, "R2": r ** 2, "confiabilidad": r ** 2 * 100, "nivel": nivel(r)}


def textos(c, tramos):
    r, r2, conf = c["r"], c["R2"], c["confiabilidad"]
    assert c["nivel"] == "Regular" and r > 0
    return {
        "R2": f"La recta explica el {fmt(conf)} % de la variación del número de infracciones a partir de los años de experiencia. "
              f"El {fmt(100 - conf)} % restante depende de otros factores o de variación que la recta no recoge.",
        "confiabilidad": f"Es el mismo R² expresado en porcentaje: {fmt(conf)} %. El modelo da una idea general de cuántas infracciones se esperan según la experiencia, "
                         f"pero deja sin explicar cerca del {fmt(100 - conf, 0)} % de la variación, así que sus predicciones son orientativas y no exactas.",
        "r": f"Es {fmt(r, 4)}, un valor positivo: al aumentar los años de experiencia tienden a aumentar las infracciones. Elevado al cuadrado da {fmt(r2, 4)}, que es el R².",
        "nivel": f"Regular según la escala adoptada (|r| entre 0,50 y 0,80). Como r = {fmt(r, 2)} está cerca del límite de 0,80 que separaría 'aceptable', "
                 f"conviene citar el valor de r junto con la etiqueta.",
        "nota": f"Estos valores resumen el ajuste global de la recta. El diagrama mostró tres tramos y, entre 2 y 10 años de experiencia, el promedio de infracciones casi no cambia "
                f"(entre {fmt(tramos['medio']['media_min'])} y {fmt(tramos['medio']['media_max'])}), así que un R² de {fmt(conf)} % no significa que las infracciones suban de forma continua con cada año. "
                f"Además, una correlación, por alta que sea, no prueba que la experiencia cause las infracciones.",
    }


def escribir_correlacion(wb, ws, df, info, valores=True):
    c = calcular(df)
    rx, ry = rango(df, X), rango(df, Y)
    tx = textos(c, info["tramos"])
    v = (lambda t: t) if valores else (lambda t: None)

    sub = wb.add_format({"bold": True, "font_size": 13})
    neg = wb.add_format({"bold": True})
    enc = wb.add_format({"bold": True, "bg_color": "#D9E2F3", "border": 1, "align": "center", "valign": "vcenter"})
    etq = wb.add_format({"bold": True, "border": 1, "valign": "vcenter", "text_wrap": True})
    n4 = wb.add_format({"border": 1, "num_format": "0.0000", "valign": "vcenter"})
    n2 = wb.add_format({"border": 1, "num_format": "0.00", "valign": "vcenter"})
    tt = wb.add_format({"border": 1, "align": "right", "bold": True, "valign": "vcenter"})
    txt = wb.add_format({"border": 1, "text_wrap": True, "valign": "vcenter"})
    chico = wb.add_format({"border": 1, "text_wrap": True, "valign": "vcenter", "font_color": "#595959", "font_size": 9})
    par = wb.add_format({"text_wrap": True, "valign": "top"})
    cen = wb.add_format({"border": 1, "align": "center"})

    f = info["fila_libre"]
    ws.write(f, 0, "Coeficiente de determinación, confiabilidad y correlación de Pearson", sub)
    f += 2
    ws.write(f, 0, "Medida", enc)
    ws.write(f, 1, "Resultado", enc)
    ws.merge_range(f, 2, f, 8, "Interpretación", enc)
    ws.merge_range(f, 9, f, 12, "Fórmula de Excel", enc)
    fr2, fr = f + 2, f + 4   # filas de Excel de R² y de r
    filas = [
        ("Coeficiente de determinación (R²)", f"=RSQ({ry},{rx})", c["R2"], n4, tx["R2"], "RSQ(y, x)"),
        ("Confiabilidad del modelo (%)", f"=B{fr2}*100", c["confiabilidad"], n2, tx["confiabilidad"], "R² x 100"),
        ("Coeficiente de correlación de Pearson (r)", f"=CORREL({ry},{rx})", c["r"], n4, tx["r"], "CORREL(y, x)"),
        ("Nivel de correlación lineal",
         f'=IF(ABS(B{fr})=1,"Perfecta",IF(ABS(B{fr})>=0.9,"Excelente",IF(ABS(B{fr})>=0.8,"Aceptable",IF(ABS(B{fr})>=0.5,"Regular",IF(ABS(B{fr})>0,"Mínima","Sin correlación")))))',
         c["nivel"], tt, tx["nivel"], "IF anidado sobre |r|"),
    ]
    celdas = {}
    for i, (nombre, formula, val, fm, texto, ftxt) in enumerate(filas, start=1):
        ws.set_row(f + i, 48)
        ws.write(f + i, 0, nombre, etq)
        ws.write_formula(f + i, 1, formula, fm, v(val))
        ws.merge_range(f + i, 2, f + i, 8, texto, txt)
        ws.merge_range(f + i, 9, f + i, 12, ftxt, chico)
        celdas[nombre] = (f + i + 1, 2)
    assert celdas["Coeficiente de correlación de Pearson (r)"][0] == fr and celdas["Coeficiente de determinación (R²)"][0] == fr2

    f += len(filas) + 2
    ws.write(f, 0, "Escala adoptada para el nivel de correlación (|r|)", neg)
    ws.write(f + 1, 0, "Intervalo de |r|", enc)
    ws.write(f + 1, 1, "Nivel", enc)
    escala = [("|r| = 1", "Perfecta"), ("0,90 a menos de 1", "Excelente"), ("0,80 a menos de 0,90", "Aceptable"),
              ("0,50 a menos de 0,80", "Regular"), ("Mayor que 0 y menor que 0,50", "Mínima"), ("|r| = 0", "Sin correlación")]
    for i, (a, b) in enumerate(escala, start=2):
        ws.write(f + i, 0, a, cen)
        ws.write(f + i, 1, b, cen)
    f += len(escala) + 3
    ws.merge_range(f, 0, f, 12, "Esta escala es una convención de uso común que adoptamos en este trabajo. No proviene de las lecturas del curso, y otras escalas usan límites distintos.", par)
    ws.set_row(f, 30)
    f += 2
    ws.write(f, 0, "Cómo leer estos resultados", neg)
    ws.merge_range(f + 1, 0, f + 1, 12, tx["nota"], par)
    ws.set_row(f + 1, 62)
    info["fila_libre_correlacion"] = f + 3
    return {"celdas": celdas, "calc": c, "textos": tx}


def construir(ruta, df, valores=True):
    wb = xlsxwriter.Workbook(str(ruta), {"use_future_functions": True})
    ws = wb.add_worksheet(HOJA)
    info = escribir_regresion(wb, ws, df, valores=valores)
    info_c = escribir_correlacion(wb, ws, df, info, valores=valores)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    escribir_datos(wb, df)
    wb.close()
    return info, info_c


def verificar(df):
    x, y = df[X].to_numpy(), df[Y].to_numpy()
    c = calcular(df)
    pear = stats.pearsonr(x, y)
    lin = stats.linregress(x, y)
    assert abs(c["r"] - pear.statistic) < 1e-12 and abs(c["r"] - lin.rvalue) < 1e-12
    assert abs(c["R2"] - lin.rvalue ** 2) < 1e-12
    print(f"numpy corrcoef r = {c['r']:.10f}; scipy pearsonr = {pear.statistic:.10f}; linregress = {lin.rvalue:.10f}; r² = {c['R2']:.10f}")
    print(f"|r| = {abs(c['r']):.4f} cae en la escala como {c['nivel']}")
    with tempfile.TemporaryDirectory() as tmp:
        sin_cache = Path(tmp) / "ej3_correlacion_sin_cache.xlsx"
        info, info_c = construir(sin_cache, df, valores=False)
        hoja = openpyxl.load_workbook(recalcular_con_libreoffice(sin_cache, tmp), data_only=True)[HOJA]
        leidos = {k: hoja.cell(*rc).value for k, rc in info_c["celdas"].items()}
    esperados = {"Coeficiente de determinación (R²)": c["R2"], "Confiabilidad del modelo (%)": c["confiabilidad"],
                 "Coeficiente de correlación de Pearson (r)": c["r"], "Nivel de correlación lineal": c["nivel"]}
    for k, esp in esperados.items():
        got = leidos[k]
        ok = got == esp if isinstance(esp, str) else abs(got - esp) < 1e-9
        print(f"  {'OK ' if ok else 'ERROR'} {k}: Excel(LibreOffice) = {got}; Python = {esp}")
        assert ok, k
    # Otros límites de la escala, para comprobar los IF anidados con valores de borde
    for r, esp in [(1, "Perfecta"), (-0.95, "Excelente"), (0.8, "Aceptable"), (-0.5, "Regular"), (0.49, "Mínima"), (0, "Sin correlación")]:
        assert nivel(r) == esp, (r, nivel(r))
    print("Función de nivel comprobada en los límites de la escala.")


if __name__ == "__main__":
    df = cargar_datos()
    salida = BASE / "output"
    salida.mkdir(exist_ok=True)
    info, info_c = construir(salida / "ej3_correlacion_borrador.xlsx", df)
    for k, t in info_c["textos"].items():
        print(f"{k}: {t}")
    print()
    verificar(df)
