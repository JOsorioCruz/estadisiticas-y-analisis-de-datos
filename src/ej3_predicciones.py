"""Ejercicio 3: tres predicciones con la ecuación de regresión."""
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import xlsxwriter
from scipy import stats

from carga import BASE, cargar_datos
from ej2_tabla import recalcular_con_libreoffice
from ej3_correlacion import escribir_correlacion
from ej3_regresion import HOJA, X, Y, escribir_regresion, fmt
from hoja_datos import escribir_datos, rango

# Un valor de experiencia por cada tramo observado (0 a 1, 2 a 10 y 11 o más años); todos dentro del rango 0 a 30
PREDICCIONES = [1, 5, 12]


def calcular(df, xs=PREDICCIONES):
    x, y = df[X].to_numpy(), df[Y].to_numpy()
    r = stats.linregress(x, y)
    filas = []
    for xv in xs:
        m = x == xv
        filas.append({"x": xv, "pred": float(r.intercept + r.slope * xv), "obs": float(y[m].mean()), "n": int(m.sum()),
                      "dentro": bool(x.min() <= xv <= x.max())})
    error = float(np.sqrt(np.sum((y - (r.intercept + r.slope * x)) ** 2) / (len(x) - 2)))   # error típico de la estimación
    return {"filas": filas, "error": error, "xmin": int(x.min()), "xmax": int(x.max())}


def textos(c):
    p1, p2, p3 = c["filas"]
    assert all(f["dentro"] and f["n"] > 0 for f in c["filas"])
    assert p1["pred"] > p1["obs"] and abs(p2["pred"] - p2["obs"]) < 0.5 and p3["pred"] < p3["obs"]
    t = []
    for f, cierre in zip(c["filas"], (
            "queda por encima de lo observado y la recta sobreestima en este tramo.",
            "coincide bien con lo observado.",
            "queda por debajo de lo observado y la recta subestima en este tramo.")):
        anios = "1 año" if f["x"] == 1 else f"{f['x']} años"
        t.append(f"Para un conductor con {anios} de experiencia, el modelo estima {fmt(f['pred'])} infracciones en los últimos cinco años. "
                 f"Los {f['n']} conductores con {anios} tuvieron en promedio {fmt(f['obs'])}; la estimación {cierre}")
    general = (f"Las tres predicciones usan valores dentro del rango observado ({c['xmin']} a {c['xmax']} años), que es donde la ecuación puede usarse. "
               f"El error típico de la estimación es de {fmt(c['error'])} infracciones, así que cada predicción debe leerse con un margen de ese orden. "
               f"Por el patrón de tres tramos visto en el diagrama, la recta funciona mejor en el tramo intermedio (2 a 10 años) que en los extremos.")
    return t, general


def escribir_predicciones(wb, ws, df, info, valores=True):
    c = calcular(df)
    rx, ry = rango(df, X), rango(df, Y)
    interp, general = textos(c)
    v = (lambda t: t) if valores else (lambda t: None)
    fb, fa = info["celdas"]["Pendiente (b)"][0], info["celdas"]["Intercepto (a)"][0]   # filas de Excel de b y a

    sub = wb.add_format({"bold": True, "font_size": 13})
    neg = wb.add_format({"bold": True})
    enc = wb.add_format({"bold": True, "bg_color": "#D9E2F3", "border": 1, "align": "center", "valign": "vcenter", "text_wrap": True})
    etq = wb.add_format({"bold": True, "border": 1, "valign": "vcenter", "text_wrap": True})
    ent = wb.add_format({"bg_color": "#FFF2CC", "border": 1, "num_format": "0", "valign": "vcenter"})
    n0 = wb.add_format({"border": 1, "num_format": "0", "valign": "vcenter"})
    n2 = wb.add_format({"border": 1, "num_format": "0.00", "valign": "vcenter"})
    ce = wb.add_format({"border": 1, "align": "center", "valign": "vcenter"})
    txt = wb.add_format({"border": 1, "text_wrap": True, "valign": "vcenter"})
    par = wb.add_format({"text_wrap": True, "valign": "top"})

    f = info["fila_libre_correlacion"]
    ws.write(f, 0, "Predicciones con la ecuación de regresión", sub)
    ws.write(f + 1, 0, "La ecuación es y = a + b x, con a y b tomados de las celdas de la pendiente y el intercepto. Las celdas amarillas se pueden cambiar.")
    f += 3
    ws.set_row(f, 44)
    for j, e in enumerate(["Predicción", "Experiencia x (años)", "Infracciones estimadas ŷ", "Promedio observado con ese x",
                           "Conductores con ese x", "¿x dentro del rango observado?"]):
        ws.write(f, j, e, enc)
    ws.merge_range(f, 6, f, 12, "Interpretación", enc)
    celdas = {}
    for i, (fila, texto) in enumerate(zip(c["filas"], interp), start=1):
        r = f + i
        x_ = r + 1   # fila de Excel
        ws.set_row(r, 62)
        ws.write(r, 0, f"Predicción {i}", etq)
        ws.write_number(r, 1, fila["x"], ent)
        ws.write_formula(r, 2, f"=$B${fa}+$B${fb}*B{x_}", n2, v(fila["pred"]))
        ws.write_formula(r, 3, f"=AVERAGEIF({rx},B{x_},{ry})", n2, v(fila["obs"]))
        ws.write_formula(r, 4, f"=COUNTIF({rx},B{x_})", n0, v(fila["n"]))
        ws.write_formula(r, 5, f'=IF(AND(B{x_}>=MIN({rx}),B{x_}<=MAX({rx})),"Sí","No")', ce, v("Sí"))
        ws.merge_range(r, 6, r, 12, texto, txt)
        celdas[f"pred{i}"] = (x_, 3)
        celdas[f"obs{i}"] = (x_, 4)
        celdas[f"n{i}"] = (x_, 5)
        celdas[f"dentro{i}"] = (x_, 6)
    f += len(c["filas"]) + 2
    ws.write(f, 0, "Error típico de la estimación (STEYX)", etq)
    ws.write_formula(f, 1, f"=STEYX({ry},{rx})", n2, v(c["error"]))
    celdas["error"] = (f + 1, 2)
    f += 2
    ws.write(f, 0, "Cómo usar las predicciones", neg)
    ws.merge_range(f + 1, 0, f + 1, 12, general, par)
    ws.set_row(f + 1, 62)
    return {"celdas": celdas, "calc": c, "textos": interp + [general]}


def construir(ruta, df, valores=True):
    wb = xlsxwriter.Workbook(str(ruta), {"use_future_functions": True})
    ws = wb.add_worksheet(HOJA)
    info = escribir_regresion(wb, ws, df, valores=valores)
    info_c = escribir_correlacion(wb, ws, df, info, valores=valores)
    info_p = escribir_predicciones(wb, ws, df, info, valores=valores)
    info_p["celdas_regresion"], info_p["celdas_correlacion"] = info["celdas"], info_c["celdas"]
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    escribir_datos(wb, df)
    wb.close()
    return info_p


def verificar(df):
    x, y = df[X].to_numpy(), df[Y].to_numpy()
    c = calcular(df)
    b, a = np.polyfit(x, y, 1)            # segunda implementación, independiente de scipy
    for f in c["filas"]:
        assert abs(f["pred"] - (a + b * f["x"])) < 1e-9
    with tempfile.TemporaryDirectory() as tmp:
        sin_cache = Path(tmp) / "ej3_pred_sin_cache.xlsx"
        info_p = construir(sin_cache, df, valores=False)
        hoja = openpyxl.load_workbook(recalcular_con_libreoffice(sin_cache, tmp), data_only=True)[HOJA]
        leidos = {k: hoja.cell(*rc).value for k, rc in info_p["celdas"].items()}
    peor = 0.0
    print(f"{'x':>3}{'ŷ Excel':>12}{'ŷ polyfit':>12}{'obs Excel':>12}{'obs numpy':>12}{'n Excel':>9}{'n numpy':>9}")
    for i, f in enumerate(c["filas"], start=1):
        e = (leidos[f"pred{i}"], leidos[f"obs{i}"], leidos[f"n{i}"])
        peor = max(peor, abs(e[0] - (a + b * f["x"])), abs(e[1] - f["obs"]), abs(e[2] - f["n"]))
        assert leidos[f"dentro{i}"] == "Sí"
        print(f"{f['x']:>3}{e[0]:>12.4f}{a + b * f['x']:>12.4f}{e[1]:>12.4f}{f['obs']:>12.4f}{e[2]:>9}{f['n']:>9}")
    err = float(np.sqrt(np.sum((y - (a + b * x)) ** 2) / (len(x) - 2)))
    peor = max(peor, abs(leidos["error"] - err), abs(leidos["error"] - c["error"]))
    print(f"STEYX Excel = {leidos['error']:.6f}; numpy = {err:.6f}")
    assert peor < 1e-9, peor
    print(f"Diferencia máxima Excel (LibreOffice) contra numpy: {peor:.2e}")


if __name__ == "__main__":
    df = cargar_datos()
    salida = BASE / "output"
    salida.mkdir(exist_ok=True)
    info_p = construir(salida / "ej3_predicciones_borrador.xlsx", df)
    for t in info_p["textos"]:
        print(t + "\n")
    verificar(df)
