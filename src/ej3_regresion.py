"""Ejercicio 3: ecuación de regresión lineal y diagrama de dispersión (experiencia como x, infracciones como y)."""
import tempfile
import zipfile
import re
from pathlib import Path

import numpy as np
import openpyxl
import xlsxwriter
from scipy import stats

from carga import BASE, cargar_datos
from ej2_tabla import recalcular_con_libreoffice
from hoja_datos import escribir_datos, rango

X = "AÑOS DE EXPERIENCIA DEL CONDUCTOR"
Y = "HISTÓRICO DE INFRACCIONES"
HOJA = "Ej3 Regresión"
COLOR = "#2F5597"


def fmt(x, d=2):
    """Número con coma decimal y d decimales."""
    return f"{x:.{d}f}".replace(".", ",")


def calcular(df):
    x, y = df[X].to_numpy(), df[Y].to_numpy()
    r = stats.linregress(x, y)
    return {
        "n": len(x), "media_x": float(x.mean()), "media_y": float(y.mean()),
        "pendiente": float(r.slope), "intercepto": float(r.intercept),
        "y_con_0": float(y[x == 0].mean()), "casos_con_0": int((x == 0).sum()),
        "tipo": "positiva" if r.slope > 0 else ("negativa" if r.slope < 0 else "sin relación"),
    }


def tramos(df):
    """Resumen de y en tres tramos de experiencia (0 a 1, 2 a 10, 11 o más), para describir el patrón de los datos."""
    x, y = df[X], df[Y]
    limites = [("bajo", x <= 1), ("medio", (x >= 2) & (x <= 10)), ("alto", x >= 11)]
    res = {}
    for nombre, m in limites:
        medias = y[m].groupby(x[m]).mean()
        res[nombre] = {"n": int(m.sum()), "min": int(y[m].min()), "max": int(y[m].max()),
                       "media_min": float(medias.min()), "media_max": float(medias.max())}
    assert sum(t["n"] for t in res.values()) == len(df)
    return res


def ecuacion_texto(a, b):
    """Igual que la fórmula de Excel: y = a + b x, con ROUND a 4 decimales."""
    return f"y = {round(a, 4)} {'+' if b >= 0 else '-'} {abs(round(b, 4))} x"


def textos(c, t):
    a, b = c["intercepto"], c["pendiente"]
    assert c["tipo"] == "positiva"
    # El texto describe tres tramos separados: se comprueba que siguen siéndolo con estos datos
    assert t["bajo"]["max"] <= t["medio"]["min"] and t["medio"]["max"] < t["alto"]["min"]
    assert t["medio"]["media_max"] - t["medio"]["media_min"] < 2
    return {
        "Pendiente": f"Por cada año más de experiencia, el modelo asocia unas {fmt(b)} infracciones adicionales en los últimos cinco años. "
                     f"Esto describe una asociación y no demuestra que la experiencia cause las infracciones.",
        "Intercepto": f"Es el valor que da la recta para un conductor con 0 años de experiencia: {fmt(a)} infracciones. "
                      f"Ojo: los {c['casos_con_0']} conductores con 0 años tuvieron en promedio {fmt(c['y_con_0'])}, así que la recta se aleja algo de los datos en ese extremo.",
        "Tipo de relación": f"Positiva. La pendiente es mayor que cero y la nube de puntos sube de izquierda a derecha: a más años de experiencia, "
                            f"tienden a ser más las infracciones acumuladas. La fuerza de esa relación se mide con el coeficiente de determinación y con r.",
        "Patrón": f"Los datos no forman una nube pareja, sino tres tramos. Con 0 o 1 año de experiencia ({t['bajo']['n']} conductores) las infracciones van de {t['bajo']['min']} a {t['bajo']['max']}. "
                  f"Con 2 a 10 años ({t['medio']['n']}) van de {t['medio']['min']} a {t['medio']['max']} y su promedio casi no cambia de un año a otro (entre {fmt(t['medio']['media_min'])} y {fmt(t['medio']['media_max'])}). "
                  f"Con 11 años o más ({t['alto']['n']}) van de {t['alto']['min']} a {t['alto']['max']}. La recta resume ese escalón con una sola pendiente, así que sirve como aproximación general "
                  f"pero no describe bien cada tramo por separado.",
        "Gráfico": "Como las dos variables son números enteros, muchos conductores comparten la misma pareja de valores y sus puntos quedan uno encima de otro. "
                   "Por eso se ven menos puntos que los 400 registros.",
    }


def escribir_regresion(wb, ws, df, valores=True):
    c = calcular(df)
    rx, ry = rango(df, X), rango(df, Y)
    v = (lambda t: t) if valores else (lambda t: None)
    tx = textos(c, tramos(df))

    titulo = wb.add_format({"bold": True, "font_size": 13})
    neg = wb.add_format({"bold": True})
    enc = wb.add_format({"bold": True, "bg_color": "#D9E2F3", "border": 1, "align": "center", "valign": "vcenter"})
    etq = wb.add_format({"bold": True, "border": 1, "valign": "vcenter"})
    n0 = wb.add_format({"border": 1, "num_format": "0", "valign": "vcenter"})
    n2 = wb.add_format({"border": 1, "num_format": "0.00", "valign": "vcenter"})
    n4 = wb.add_format({"border": 1, "num_format": "0.0000", "valign": "vcenter"})
    tt = wb.add_format({"border": 1, "align": "right", "bold": True, "valign": "vcenter"})
    txt = wb.add_format({"border": 1, "text_wrap": True, "valign": "vcenter"})
    chico = wb.add_format({"border": 1, "text_wrap": True, "valign": "vcenter", "font_color": "#595959", "font_size": 9})
    par = wb.add_format({"text_wrap": True, "valign": "top"})

    ws.write(0, 0, "Regresión lineal: años de experiencia (x) e histórico de infracciones (y)", titulo)
    ws.write(1, 0, "Variable independiente (x): años de experiencia del conductor. Variable dependiente (y): número de infracciones en los últimos cinco años.")
    ws.set_column(0, 0, 40)
    ws.set_column(1, 1, 24)
    ws.set_column(2, 8, 12)
    ws.set_column(9, 12, 12)

    # Datos de partida
    ws.write(3, 0, "Datos de partida", neg)
    ws.write(4, 0, "Número de pares de datos (n)", etq)
    ws.write_formula(4, 1, f"=COUNT({rx})", n0, v(c["n"]))
    ws.write(5, 0, "Media de x (años de experiencia)", etq)
    ws.write_formula(5, 1, f"=AVERAGE({rx})", n2, v(c["media_x"]))
    ws.write(6, 0, "Media de y (infracciones)", etq)
    ws.write_formula(6, 1, f"=AVERAGE({ry})", n2, v(c["media_y"]))

    # Ecuación
    ws.write(8, 0, "Ecuación de regresión lineal (y = a + b x)", neg)
    ws.write(9, 0, "Medida", enc)
    ws.write(9, 1, "Resultado", enc)
    ws.merge_range(9, 2, 9, 8, "Interpretación", enc)
    ws.merge_range(9, 9, 9, 12, "Fórmula de Excel", enc)
    filas = [
        ("Pendiente (b)", f"=SLOPE({ry},{rx})", c["pendiente"], n4, tx["Pendiente"], "SLOPE(y, x)"),
        ("Intercepto (a)", f"=INTERCEPT({ry},{rx})", c["intercepto"], n4, tx["Intercepto"], "INTERCEPT(y, x)"),
    ]
    celdas = {}
    for i, (nombre, formula, val, fm, texto, ftxt) in enumerate(filas, start=10):
        ws.set_row(i, 48)
        ws.write(i, 0, nombre, etq)
        ws.write_formula(i, 1, formula, fm, v(val))
        ws.merge_range(i, 2, i, 8, texto, txt)
        ws.merge_range(i, 9, i, 12, ftxt, chico)
        celdas[nombre] = (i + 1, 2)
    fb, fa = 11, 12   # filas de Excel de b y de a
    ws.set_row(12, 34)
    ws.write(12, 0, "Ecuación de la recta", etq)
    eq_formula = f'="y = "&ROUND(B{fa},4)&IF(B{fb}>=0," + "," - ")&ABS(ROUND(B{fb},4))&" x"'
    ws.write_formula(12, 1, eq_formula, tt, v(ecuacion_texto(c["intercepto"], c["pendiente"])))
    ws.merge_range(12, 2, 12, 8, "Es la recta de mínimos cuadrados: con ella se estima el número de infracciones a partir de los años de experiencia.", txt)
    ws.merge_range(12, 9, 12, 12, "\"y = \" & a & \" + \" & b & \" x\"", chico)
    celdas["Ecuación"] = (13, 2)
    ws.set_row(13, 34)
    ws.write(13, 0, "Tipo de relación según la pendiente", etq)
    ws.write_formula(13, 1, f'=IF(B{fb}>0,"Positiva",IF(B{fb}<0,"Negativa","Sin relación"))', tt, v(c["tipo"].capitalize()))
    ws.merge_range(13, 2, 13, 8, tx["Tipo de relación"], txt)
    ws.merge_range(13, 9, 13, 12, "Signo de b", chico)
    celdas["Tipo de relación"] = (14, 2)
    ws.set_row(14, 34)
    ws.write(14, 0, "Promedio observado de y con x = 0", etq)
    ws.write_formula(14, 1, f"=AVERAGEIF({rx},0,{ry})", n2, v(c["y_con_0"]))
    ws.merge_range(14, 2, 14, 8, "Se usa para comparar el intercepto de la recta con lo que ocurre realmente en los conductores sin experiencia.", txt)
    ws.merge_range(14, 9, 14, 12, "AVERAGEIF(x, 0, y)", chico)
    celdas["Promedio con x = 0"] = (15, 2)

    # Diagrama de dispersión con la recta de regresión
    graf = wb.add_chart({"type": "scatter"})
    graf.add_series({
        "name": "Conductores (un punto por accidente)",
        "categories": f"=DATOS!{rx.split('!')[1]}",
        "values": f"=DATOS!{ry.split('!')[1]}",
        "marker": {"type": "circle", "size": 6, "fill": {"color": COLOR, "transparency": 65}, "border": {"color": COLOR, "transparency": 30}},
        "trendline": {"type": "linear", "name": "Recta de regresión", "display_equation": True,
                      "line": {"color": "#C00000", "width": 2.25}},
    })
    graf.set_title({"name": "Diagrama de dispersión: años de experiencia e infracciones, con recta de regresión", "name_font": {"size": 12}})
    graf.set_x_axis({"name": "Años de experiencia del conductor (años)", "min": 0, "max": 30, "major_unit": 5})
    graf.set_y_axis({"name": "Histórico de infracciones (número en los últimos 5 años)", "min": 0, "max": 20, "major_unit": 2,
                     "major_gridlines": {"visible": True, "line": {"color": "#D9D9D9"}}})
    graf.set_legend({"position": "bottom"})
    graf.set_size({"width": 760, "height": 440})
    ws.write(16, 0, "Diagrama de dispersión con la recta de regresión", titulo)
    ws.insert_chart(17, 0, graf)
    ws.write(39, 0, "Lectura del diagrama", neg)
    ws.merge_range(40, 0, 40, 12, tx["Tipo de relación"] + " " + tx["Gráfico"], par)
    ws.set_row(40, 48)
    ws.write(42, 0, "Patrón de los datos", neg)
    ws.merge_range(43, 0, 43, 12, tx["Patrón"], par)
    ws.set_row(43, 62)
    return {"celdas": celdas, "calc": c, "textos": tx, "tramos": tramos(df), "fila_libre": 46}


def construir(ruta, df, valores=True):
    wb = xlsxwriter.Workbook(str(ruta), {"use_future_functions": True})
    ws = wb.add_worksheet(HOJA)
    info = escribir_regresion(wb, ws, df, valores=valores)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    escribir_datos(wb, df)
    wb.close()
    return info


def verificar(df, ruta):
    c = calcular(df)
    x, y = df[X].to_numpy(), df[Y].to_numpy()
    b_np, a_np = np.polyfit(x, y, 1)      # segunda implementación independiente
    print(f"scipy.linregress: b = {c['pendiente']:.10f}, a = {c['intercepto']:.10f}; numpy.polyfit: b = {b_np:.10f}, a = {a_np:.10f}")
    assert abs(b_np - c["pendiente"]) < 1e-9 and abs(a_np - c["intercepto"]) < 1e-9
    with tempfile.TemporaryDirectory() as tmp:
        sin_cache = Path(tmp) / "ej3_sin_cache.xlsx"
        info = construir(sin_cache, df, valores=False)
        hoja = openpyxl.load_workbook(recalcular_con_libreoffice(sin_cache, tmp), data_only=True)[HOJA]
        leidos = {k: hoja.cell(*rc).value for k, rc in info["celdas"].items()}
    esperados = {"Pendiente (b)": c["pendiente"], "Intercepto (a)": c["intercepto"],
                 "Ecuación": ecuacion_texto(c["intercepto"], c["pendiente"]), "Tipo de relación": "Positiva",
                 "Promedio con x = 0": c["y_con_0"]}
    for k, esp in esperados.items():
        got = leidos[k]
        ok = got == esp if isinstance(esp, str) else abs(got - esp) < 1e-9
        print(f"  {'OK ' if ok else 'ERROR'} {k}: Excel(LibreOffice) = {got}; Python = {esp}")
        assert ok, k
    # Gráfico nativo, con línea de tendencia lineal y sin imágenes
    with zipfile.ZipFile(ruta) as z:
        graficos = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        xml = z.read(graficos[0]).decode("utf-8")
        print("Gráficos:", graficos, "| imágenes:", [n for n in z.namelist() if n.startswith("xl/media/")] or "ninguna")
        print("scatterChart:", "scatterChart" in xml, "| trendline lineal:", '<c:trendlineType val="linear"/>' in xml,
              "| ecuación visible:", '<c:dispEq val="1"/>' in xml)
        print("Referencias:", sorted(set(re.findall(r"<c:f>([^<]+)</c:f>", xml))))


if __name__ == "__main__":
    df = cargar_datos()
    salida = BASE / "output"
    salida.mkdir(exist_ok=True)
    ruta = salida / "ej3_regresion_borrador.xlsx"
    info = construir(ruta, df)
    print(ecuacion_texto(info["calc"]["intercepto"], info["calc"]["pendiente"]))
    for k, t in info["textos"].items():
        print(f"{k}: {t}")
    print()
    verificar(df, ruta)
