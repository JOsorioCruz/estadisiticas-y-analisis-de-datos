"""Tarea 14: genera el libro de Excel completo de la Actividad No 2 (AF17401).

Se ejecuta desde la raíz con `.venv/bin/python src/generar_libro.py`. Es reproducible: no usa rutas absolutas ni pasos manuales.
"""
import json
import math
import re
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import xlsxwriter

from carga import BASE, RUTA_ANEXO, cargar_datos
from ej1_tabla import RUTULO_QUIZ, TABLA_1, comprobar_cifras
from ej2_dispersion import escribir_dispersion
from ej2_graficos import escribir_graficos
from ej2_medidas import escribir_medidas
from ej2_preguntas import escribir_preguntas
from ej2_tabla import VARIABLE, calcular_tabla, escribir_tabla, recalcular_con_libreoffice
from ej3_correlacion import escribir_correlacion
from ej3_predicciones import PREDICCIONES, escribir_predicciones
from ej3_regresion import escribir_regresion
from hoja_datos import escribir_datos
from textos_libro import PENDIENTE_FECHA, PENDIENTE_TUTOR, sin_frase_quiz, textos
from verificacion import comparar, esperados_ej2, esperados_ej3

NOMBRE_ARCHIVO = "actividad_probabilidad_y_estadistica_unidad_2.xlsx"   # por confirmar con Jairo
RUTA_LIBRO = BASE / "output" / NOMBRE_ARCHIVO
URL_REPOSITORIO = "https://github.com/JOsorioCruz/estadisiticas-y-analisis-de-datos"   # solo se imprime en la portada del PDF
FILA0_EJ2 = 3          # la hoja del ejercicio 2 deja tres filas arriba para el título y la variable elegida
HOJAS = ["Portada", "Presentación", "Introducción", "Objetivos", "Ejercicio 1", "Ejercicio 2", "Ejercicio 3",
         "Conclusiones", "Bibliografía", "Anexos", "DATOS"]
ERRORES_EXCEL = ("#NAME?", "#VALUE!", "#DIV/0!", "#REF!", "#N/A", "#NUM!", "#NULL!")


class Estilos:
    def __init__(self, wb):
        self.titulo = wb.add_format({"bold": True, "font_size": 16, "font_color": "#1F3864"})
        self.subtitulo = wb.add_format({"bold": True, "font_size": 12, "font_color": "#1F3864"})
        self.parrafo = wb.add_format({"text_wrap": True, "valign": "top"})
        self.item = wb.add_format({"text_wrap": True, "valign": "top", "indent": 1})
        self.enc = wb.add_format({"bold": True, "bg_color": "#D9E2F3", "border": 1, "align": "center", "valign": "vcenter", "text_wrap": True})
        self.celda = wb.add_format({"border": 1, "text_wrap": True, "valign": "top"})
        self.celda_neg = wb.add_format({"border": 1, "bold": True, "text_wrap": True, "valign": "top"})
        self.etq = wb.add_format({"bold": True, "font_size": 12, "valign": "top"})
        self.valor = wb.add_format({"font_size": 12, "valign": "top", "text_wrap": True})
        self.pendiente = wb.add_format({"font_size": 12, "valign": "top", "bg_color": "#FFF2CC", "bold": True})
        self.cursiva = wb.add_format({"italic": True})
        self.num = wb.add_format({"border": 1, "num_format": "0.000000", "valign": "top"})
        self.dif = wb.add_format({"border": 1, "num_format": "0.0E+00", "valign": "top"})
        self.centro = wb.add_format({"border": 1, "align": "center", "valign": "top"})
        self.quiz = wb.add_format({"border": 2, "border_color": "#7F7F7F", "align": "center", "valign": "vcenter", "italic": True,
                                   "font_color": "#7F7F7F", "font_size": 12, "bg_color": "#F2F2F2"})


def altura(texto, ancho):
    """Altura de fila (puntos) para un texto envuelto en celdas combinadas de ancho total `ancho` (en caracteres)."""
    lineas = sum(max(1, math.ceil(len(p) / (ancho * 1.15))) for p in str(texto).split("\n"))
    return 15 * lineas + 3


def escribir_parrafo(ws, fila, texto, fmt, hasta_col, ancho):
    if hasta_col == 0:
        ws.write(fila, 0, texto, fmt)      # una sola columna: no hay nada que combinar
    else:
        ws.merge_range(fila, 0, fila, hasta_col, texto, fmt)
    ws.set_row(fila, altura(texto, ancho))


def escribir_referencia(wb, ws, fila, texto, hasta_col, ancho, est):
    """Referencia APA con cursiva: el texto entre asteriscos va en cursiva."""
    partes = [p for p in re.split(r"\*", texto)]
    args = []
    for i, p in enumerate(partes):
        if p == "":
            continue
        if i % 2 == 1:
            args.extend([est.cursiva, p])
        else:
            args.append(p)
    ws.write_rich_string(fila, 0, *args, est.parrafo)
    ws.set_row(fila, altura(texto.replace("*", ""), ancho))


def hoja_texto(wb, est, nombre, titulo, bloques, ancho=110):
    """Hoja de texto: bloques = lista de ('sub', texto) o ('p', texto) o ('li', texto)."""
    ws = wb.add_worksheet(nombre)
    ws.hide_gridlines(2)
    ws.set_column(0, 0, ancho)
    ws.write(0, 0, titulo, est.titulo)
    ws.set_row(0, 26)
    fila = 2
    for tipo, texto in bloques:
        if tipo == "sub":
            ws.write(fila, 0, texto, est.subtitulo)
            fila += 1
        else:
            escribir_parrafo(ws, fila, texto, est.item if tipo == "li" else est.parrafo, 0, ancho)
            fila += 2 if tipo == "p" else 1
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    return ws


def hoja_portada(wb, est, t, para_pdf=False):
    ws = wb.add_worksheet("Portada")
    ws.hide_gridlines(2)
    ws.set_column(0, 0, 24)
    ws.set_column(1, 1, 78)
    ws.write(1, 0, "Universidad de Cartagena", wb.add_format({"bold": True, "font_size": 20, "font_color": "#1F3864"}))
    if para_pdf:          # en el PDF, el enlace al repositorio va al inicio de la portada
        ws.write_url(0, 0, URL_REPOSITORIO, wb.add_format({"font_color": "#0563C1", "underline": 1, "font_size": 11}), string=f"Repositorio del trabajo: {URL_REPOSITORIO}")
    if not para_pdf:      # en el PDF no se imprime la línea del CTEV
        ws.write(2, 0, "Centro Tecnológico para la Formación Virtual y a Distancia (CTEV)", wb.add_format({"font_size": 12, "font_color": "#595959"}))
    fila = 5
    for etiqueta, valor in t["portada"][2:]:
        if para_pdf and etiqueta in ("Tutor", "Fecha"):      # en el PDF no se imprimen el tutor ni la fecha
            continue
        ws.write(fila, 0, etiqueta, est.etq)
        ws.write(fila, 1, valor, est.pendiente if valor.startswith("PENDIENTE") else est.valor)
        ws.set_row(fila, 24)
        fila += 1
    ws.set_landscape()
    ws.fit_to_pages(1, 1)
    return ws


def hoja_ejercicio1(wb, est, para_pdf=False):
    ws = wb.add_worksheet("Ejercicio 1")
    ws.hide_gridlines(2)
    ancho = {0: 24, 1: 72, 2: 72}
    for c, a in ancho.items():
        ws.set_column(c, c, a)
    ws.write(0, 0, "Ejercicio 1: definiciones de conceptos estadísticos", est.titulo)
    ws.set_row(0, 26)
    ws.write(2, 0, "Tabla 1. Definiciones de conceptos", est.subtitulo)
    for j, e in enumerate(["Conceptos", "Definiciones", "Ejemplos"]):
        ws.write(3, j, e, est.enc)
    for i, (concepto, definicion, ejemplo) in enumerate(TABLA_1, start=4):
        ws.write(i, 0, concepto, est.celda_neg)
        ws.write(i, 1, definicion, est.celda)
        ws.write(i, 2, ejemplo, est.celda)
        ws.set_row(i, max(altura(definicion, ancho[1]), altura(ejemplo, ancho[2])))
    if not para_pdf:      # en el PDF no se imprime el recuadro de la evidencia del quiz
        fila = 4 + len(TABLA_1) + 1
        ws.write(fila, 0, "Evidencia del quiz", est.subtitulo)
        ws.merge_range(fila + 1, 0, fila + 12, 2, RUTULO_QUIZ, est.quiz)
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    return ws


def configurar_impresion(ws, saltos, ultima_fila=None, ultima_col=12):
    """Hoja ancha en A4 horizontal con escala fija y saltos de página antes de cada bloque, para que ningún gráfico quede partido."""
    ws.set_landscape()
    ws.set_paper(9)
    ws.set_margins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.set_print_scale(60)
    ws.set_h_pagebreaks(sorted(saltos))
    if ultima_fila is not None:
        ws.print_area(0, 0, ultima_fila, ultima_col)      # evita páginas en blanco por filas vacías al final


def hoja_ejercicio2(wb, df, valores):
    ws = wb.add_worksheet("Ejercicio 2")
    ws.write(0, 0, "Ejercicio 2: estadística descriptiva de la velocidad registrada", wb.add_format({"bold": True, "font_size": 16, "font_color": "#1F3864"}))
    ws.set_row(0, 26)
    ws.write(1, 0, "Variable elegida: velocidad registrada (km/h), variable cuantitativa continua (columna E de la hoja DATOS, 400 registros).")
    pos = escribir_tabla(wb, ws, df, valores=valores, fila0=FILA0_EJ2)
    tab = calcular_tabla(df[VARIABLE])
    escribir_graficos(wb, ws, pos, tab, valores=valores)
    info_q = escribir_preguntas(wb, ws, pos, tab, valores=valores)
    info_m = escribir_medidas(wb, ws, pos, df, valores=valores)
    info_d = escribir_dispersion(wb, ws, pos, df, info_m, valores=valores)
    # páginas: tabla | histograma y circular | polígono y preguntas | medidas | dispersión y análisis
    ic = pos["inicios_graficos"]
    configurar_impresion(ws, [ic[0] - 1, ic[2], pos["fila_libre_preguntas"], pos["fila_libre_medidas"]], pos["fila_libre_dispersion"] - 2)
    return {"pos": pos, "q": info_q, "m": info_m, "d": info_d}


def hoja_ejercicio3(wb, df, valores):
    ws = wb.add_worksheet("Ejercicio 3")
    info_r = escribir_regresion(wb, ws, df, valores=valores)
    info_c = escribir_correlacion(wb, ws, df, info_r, valores=valores)
    info_p = escribir_predicciones(wb, ws, df, info_r, valores=valores)
    # páginas: ecuación | diagrama de dispersión | R², r y nivel | predicciones (el diagrama empieza en la fila 17 de Excel)
    configurar_impresion(ws, [16, info_r["fila_libre"], info_r["fila_libre_correlacion"]])
    return {"r": info_r, "c": info_c, "p": info_p}


def hoja_anexos(wb, est, t, filas_verif, para_pdf=False):
    ws = wb.add_worksheet("Anexos")
    ws.hide_gridlines(2)
    anchos = [16, 36, 48, 22, 22, 14, 12]
    for c, a in enumerate(anchos):
        ws.set_column(c, c, a)
    total = sum(anchos)
    ws.write(0, 0, "Anexos", est.titulo)
    ws.set_row(0, 26)
    fila = 2
    an = t["anexos"]
    # Anexo A
    ws.write(fila, 0, "Anexo A. Hoja de datos", est.subtitulo)
    fila += 1
    for p in an["Anexo A. Hoja de datos"]:
        escribir_parrafo(ws, fila, p, est.parrafo, 6, total)
        fila += 1
    fila += 1
    # Anexo B
    ws.write(fila, 0, "Anexo B. Verificación de los cálculos", est.subtitulo)
    fila += 1
    for p in an["Anexo B. Verificación de los cálculos"]:
        escribir_parrafo(ws, fila, p, est.parrafo, 6, total)
        fila += 1
    fila += 1
    for j, e in enumerate(["Ejercicio", "Bloque", "Medida", "Valor en Excel (recalculado)", "Valor en Python", "Diferencia", "Resultado"]):
        ws.write(fila, j, e, est.enc)
    ws.set_row(fila, 32)
    for f in filas_verif:
        fila += 1
        ws.write(fila, 0, f["ejercicio"], est.celda)
        ws.write(fila, 1, f["bloque"], est.celda)
        ws.write(fila, 2, f["medida"], est.celda)
        for col, clave in ((3, "excel_libreoffice"), (4, "python")):
            val = f[clave]
            if isinstance(val, str):
                ws.write_string(fila, col, val, est.celda)
            else:
                ws.write_number(fila, col, float(val), est.num)
        if f["diferencia"] is None:
            ws.write(fila, 5, "", est.celda)
        else:
            ws.write_number(fila, 5, float(f["diferencia"]), est.dif)
        ws.write(fila, 6, "Coincide" if f["ok"] else "Difiere", est.centro)
    fila += 3
    # Anexo C
    ws.write(fila, 0, "Anexo C. Variables elegidas en el foro", est.subtitulo)
    fila += 1
    for j, e in enumerate(["Nombre", "Variable del ejercicio 2", "Variables del ejercicio 3"]):
        ws.write(fila, j, e, est.enc)
    ws.set_row(fila, 22)
    fila += 1
    ws.write(fila, 0, "Jairo Alonso Osorio Cruz", est.celda)
    ws.write(fila, 1, "Velocidad registrada", est.celda)
    ws.write(fila, 2, "Años de experiencia del conductor (x) e histórico de infracciones (y)", est.celda)
    ws.set_row(fila, 34)
    fila += 1
    if not para_pdf:      # en el PDF no se imprime la nota de confirmación pendiente del foro
        escribir_parrafo(ws, fila, an["Anexo C. Variables elegidas en el foro"][1], wb.add_format({"text_wrap": True, "valign": "top", "bg_color": "#FFF2CC"}), 6, total)
    fila += 3
    # Anexo D
    ws.write(fila, 0, "Anexo D. Convenciones adoptadas", est.subtitulo)
    fila += 1
    for p in an["Anexo D. Convenciones adoptadas"]:
        escribir_parrafo(ws, fila, p, est.item, 6, total)
        fila += 1
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    return ws


def construir_libro(ruta, df, valores=True, filas_verif=(), ocultar_datos=False, para_pdf=False):
    t = sin_frase_quiz(textos()) if para_pdf else textos()
    wb = xlsxwriter.Workbook(str(ruta), {"use_future_functions": True})
    wb.set_properties({"title": "Actividad No 2 de Probabilidad y Estadística (AF17401)", "subject": "Estadística descriptiva: accidentalidad en el Tolima",
                       "author": "Jairo Alonso Osorio Cruz", "comments": "Universidad de Cartagena, Ingeniería de Software"})
    est = Estilos(wb)
    hoja_portada(wb, est, t, para_pdf)
    hoja_texto(wb, est, "Presentación", "Presentación", [("p", x) for x in t["presentacion"]])
    hoja_texto(wb, est, "Introducción", "Introducción", [("p", x) for x in t["introduccion"]])
    o = t["objetivos"]
    hoja_texto(wb, est, "Objetivos", "Objetivos",
               [("sub", "Objetivo general"), ("p", o["general"]), ("sub", "Objetivos específicos")]
               + [("li", f"{i}. {x}") for i, x in enumerate(o["especificos"], 1)])
    hoja_ejercicio1(wb, est, para_pdf)
    info2 = hoja_ejercicio2(wb, df, valores)
    info3 = hoja_ejercicio3(wb, df, valores)
    bloques = []
    for titulo, texto in t["conclusiones"]:
        bloques += [("sub", titulo), ("p", texto)]
    hoja_texto(wb, est, "Conclusiones", "Conclusiones", bloques)
    # Bibliografía: solo las dos lecturas del enunciado. La lista opcional entra únicamente si Jairo la autoriza
    ws_b = wb.add_worksheet("Bibliografía")
    ws_b.hide_gridlines(2)
    ws_b.set_column(0, 0, 110)
    ws_b.write(0, 0, "Bibliografía", est.titulo)
    ws_b.set_row(0, 26)
    fila = 2
    for ref in t["bibliografia"]["Obligatoria (referencias del enunciado; confirmar que se consultaron)"]:
        escribir_referencia(wb, ws_b, fila, ref, 0, 110, est)
        fila += 2
    ws_b.set_landscape()
    ws_b.fit_to_pages(1, 0)
    hoja_anexos(wb, est, t, filas_verif, para_pdf)
    ws_d = escribir_datos(wb, df)
    ws_d.set_column(0, 0, 26)
    ws_d.set_column(1, 8, 20)
    ws_d.freeze_panes(1, 0)
    ws_d.set_tab_color("#7F7F7F")
    if ocultar_datos:
        ws_d.hide()      # solo para exportar a PDF: la hoja de 400 filas no se imprime
    wb.close()
    return {"ej2": info2, "ej3": info3}


def comparar_todo(df, info, hoja2_lo, hoja2_gu, hoja3_lo, hoja3_gu):
    """Compara cada valor de los ejercicios 2 y 3 (recalculado y guardado) con numpy y scipy."""
    filas = []
    tabla, params, preguntas, medidas, dispersion = esperados_ej2(df)
    pos = info["ej2"]["pos"]
    n_int = len(tabla["fi"])
    celdas_t = {(nom, i): (pos["primera"] + 1 + i, 2 + j) for j, nom in enumerate(tabla) for i in range(n_int)}
    esp_t = {(nom, i): tabla[nom][i] for nom in tabla for i in range(n_int)}
    nombres_t = {(nom, i): f"Tabla, intervalo {i + 1}: {nom}" for nom in tabla for i in range(n_int)}
    comparar(filas, "Ejercicio 2", "Tabla de frecuencia", esp_t, hoja2_lo, hoja2_gu, celdas_t, nombres_t)
    celdas_p = {nom: (FILA0_EJ2 + 3 + j, 2) for j, nom in enumerate(params)}
    comparar(filas, "Ejercicio 2", "Parámetros de la tabla", params, hoja2_lo, hoja2_gu, celdas_p, {k: f"Parámetro {k}" for k in params})
    comparar(filas, "Ejercicio 2", "Preguntas", preguntas, hoja2_lo, hoja2_gu, info["ej2"]["q"]["celdas"])
    comparar(filas, "Ejercicio 2", "Tendencia central y posición", medidas, hoja2_lo, hoja2_gu, info["ej2"]["m"]["celdas"])
    comparar(filas, "Ejercicio 2", "Dispersión, asimetría y curtosis", dispersion, hoja2_lo, hoja2_gu, info["ej2"]["d"]["celdas"])
    reg, cor, pred = esperados_ej3(df, PREDICCIONES)
    comparar(filas, "Ejercicio 3", "Regresión lineal", reg, hoja3_lo, hoja3_gu, info["ej3"]["r"]["celdas"])
    comparar(filas, "Ejercicio 3", "R², r y nivel", cor, hoja3_lo, hoja3_gu, info["ej3"]["c"]["celdas"])
    nombres_p = {f"{p}{i}": f"Predicción {i}: {t}" for i in range(1, 4)
                 for p, t in (("pred", "ŷ"), ("obs", "promedio observado"), ("n", "conductores"), ("dentro", "dentro del rango"))}
    nombres_p["error"] = "Error típico de la estimación"
    comparar(filas, "Ejercicio 3", "Predicciones", pred, hoja3_lo, hoja3_gu, info["ej3"]["p"]["celdas"], nombres_p)
    return filas


def hojas_de(ruta, data_only=True):
    wb = openpyxl.load_workbook(ruta, data_only=data_only)
    return wb


def calcular_verificacion(df):
    """Pasada 1: copias sin y con valores guardados, para calcular la tabla de verificación sobre el propio libro."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        info = construir_libro(tmp / "sin.xlsx", df, valores=False)
        construir_libro(tmp / "con.xlsx", df, valores=True)
        lo = openpyxl.load_workbook(recalcular_con_libreoffice(tmp / "sin.xlsx", tmp), data_only=True)
        gu = openpyxl.load_workbook(tmp / "con.xlsx", data_only=True)
        return comparar_todo(df, info, lo["Ejercicio 2"], gu["Ejercicio 2"], lo["Ejercicio 3"], gu["Ejercicio 3"])


def main():
    df = cargar_datos()
    comprobar_cifras(df)
    SALIDA = RUTA_LIBRO
    SALIDA.parent.mkdir(exist_ok=True)
    filas = calcular_verificacion(df)
    # Pasada 2: libro final con la tabla de verificación en Anexos
    info = construir_libro(SALIDA, df, valores=True, filas_verif=filas)
    print(f"Libro generado: {SALIDA.relative_to(BASE)}")
    verificar_final(SALIDA, df, info, filas)


def verificar_final(ruta, df, info, filas_previas):
    """Repite la verificación sobre el archivo final y revisa su estructura."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        lo = openpyxl.load_workbook(recalcular_con_libreoffice(ruta, tmp), data_only=True)
        gu = openpyxl.load_workbook(ruta, data_only=True)
        filas = comparar_todo(df, info, lo["Ejercicio 2"], gu["Ejercicio 2"], lo["Ejercicio 3"], gu["Ejercicio 3"])
        # Errores de Excel en cualquier celda del libro recalculado
        errores = [(h.title, c.coordinate, c.value) for h in lo.worksheets for fila in h.iter_rows() for c in fila
                   if isinstance(c.value, str) and c.value.strip() in ERRORES_EXCEL]
    n_ok = sum(f["ok"] for f in filas)
    n_gu = sum(f["valor_guardado_ok"] for f in filas)
    difs = [f["diferencia"] for f in filas if f["diferencia"] is not None]
    print(f"Verificación del archivo final: {len(filas)} comparaciones; recalculadas OK: {n_ok}; valores guardados OK: {n_gu}; dif. máx. {max(difs):.2e}")
    assert n_ok == n_gu == len(filas) == len(filas_previas), (n_ok, n_gu, len(filas), len(filas_previas))
    assert not errores, errores

    # Estructura, gráficos, imágenes y guiones largos
    wb = openpyxl.load_workbook(ruta)
    assert wb.sheetnames == HOJAS, wb.sheetnames
    print("Hojas, en orden:", " | ".join(wb.sheetnames))
    n_formulas = sum(1 for h in wb.worksheets for fila in h.iter_rows() for c in fila if isinstance(c.value, str) and c.value.startswith("="))
    print(f"Celdas con fórmula: {n_formulas}; errores de Excel tras recalcular: {len(errores)}")
    with zipfile.ZipFile(ruta) as z:
        nombres = z.namelist()
        graficos = sorted(n for n in nombres if re.match(r"xl/charts/chart\d+\.xml$", n))
        tipos = [next(t for t in ("barChart", "pieChart", "scatterChart") if t in z.read(g).decode("utf-8")) for g in graficos]
        assert not [n for n in nombres if n.startswith("xl/media/")], "hay imágenes"
        assert len(graficos) == 4 and sorted(tipos) == ["barChart", "pieChart", "scatterChart", "scatterChart"], (graficos, tipos)
        print("Gráficos nativos:", len(graficos), tipos, "| imágenes incrustadas: ninguna")
        crudo = "".join(z.read(n).decode("utf-8", "ignore") for n in nombres if n.endswith(".xml"))
        assert chr(0x2014) not in crudo, "hay un guion largo en el archivo"
        assert "_xlfn.MODE.SNGL" in crudo and "_xlfn.STDEV.S" in crudo
    print("Guiones largos (raya) en celdas, gráficos y hojas: 0")
    # Hoja DATOS idéntica al Anexo 1
    anexo = list(openpyxl.load_workbook(RUTA_ANEXO)["DATOS"].iter_rows(values_only=True))
    libro = list(wb["DATOS"].iter_rows(values_only=True))
    assert anexo == libro
    print(f"Hoja DATOS idéntica al Anexo 1: True ({len(libro) - 1} registros)")
    # Espacio del quiz y celdas pendientes
    assert any(c.value == RUTULO_QUIZ for fila in wb["Ejercicio 1"].iter_rows() for c in fila)
    pend = [(h.title, c.coordinate, c.value) for h in wb.worksheets if h.title != "DATOS" for fila in h.iter_rows() for c in fila
            if isinstance(c.value, str) and "PENDIENTE" in c.value]
    print("Espacio rotulado del quiz: presente")
    print("Campos PENDIENTE que Jairo debe completar:", pend)
    return filas


if __name__ == "__main__":
    main()
