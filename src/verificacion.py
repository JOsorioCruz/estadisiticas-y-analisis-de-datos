"""Tarea 12: verificación cruzada de cada valor del libro contra numpy y scipy, con recálculo de LibreOffice.

Los valores esperados se calculan aquí desde los datos crudos, sin reutilizar las funciones de cada módulo.
"""
import json
import tempfile
from collections import Counter
from pathlib import Path

import numpy as np
import openpyxl
from scipy import stats

from carga import BASE, RUTA_ANEXO, cargar_datos
from ej2_dispersion import construir as construir_ej2
from ej2_tabla import VARIABLE, recalcular_con_libreoffice
from ej3_predicciones import construir as construir_ej3
from ej3_regresion import X, Y

TOL = 1e-9
SALIDA = BASE / "output" / "verificacion.json"


def esperados_ej2(df):
    v = df[VARIABLE].to_numpy()
    n = len(v)
    k = int(np.ceil(1 + np.log2(n)))
    li = 32 + 11 * np.arange(k)
    fi, _ = np.histogram(v, bins=np.append(li, li[-1] + 11))
    Fi = np.cumsum(fi)
    tabla = {"Li": li, "Ls": li + 11, "xi": li + 5.5, "fi": fi, "Fi": Fi, "hi": fi / n, "Hi": Fi / n, "hi %": fi / n * 100, "Hi %": Fi / n * 100}
    imax = int(np.argmax(fi))
    m = int((Fi / n < 0.6).sum())
    etiqueta = lambda i: f"[{li[i]:g} - {li[i] + 11:g}{']' if i == k - 1 else ')'}"
    hi_a, hi_b = Fi[m - 1] / n, Fi[m] / n
    conteo = Counter(v.tolist())
    tope = max(conteo.values())
    moda = next(x for x in v.tolist() if conteo[x] == tope)
    s = v.std(ddof=1)
    params = {"n": n, "mín": v.min(), "máx": v.max(), "R": v.max() - v.min(), "k": k, "R/k": (v.max() - v.min()) / k}
    preguntas = {
        "Intervalo de mayor frecuencia": etiqueta(imax), "Frecuencia absoluta máxima (fi)": int(fi[imax]),
        "Porcentaje de los accidentes (hi %)": fi[imax] / n * 100, "Intervalo donde se alcanza el 60 %": etiqueta(m),
        "Hi (%) del intervalo anterior": hi_a * 100, "Hi (%) del intervalo donde se alcanza": hi_b * 100,
        "Estimación dentro del intervalo (km/h)": li[m] + (0.6 - hi_a) / (hi_b - hi_a) * 11,
    }
    medidas = {"Moda": moda, "Media": v.mean(), "Mediana": np.median(v), "Mínimo": v.min(), "Máximo": v.max(), "Rango": v.max() - v.min(),
               "Cuartil 1 (Q1)": np.percentile(v, 25), "Decil 5 (D5)": np.percentile(v, 50), "Cuartil 3 (Q3)": np.percentile(v, 75),
               "Percentil 10 (P10)": np.percentile(v, 10), "Percentil 80 (P80)": np.percentile(v, 80)}
    dispersion = {"Varianza muestral": v.var(ddof=1), "Desviación típica muestral": s, "Coeficiente de variación (%)": s / v.mean() * 100,
                  "Asimetría": stats.skew(v, bias=False), "Curtosis (exceso)": stats.kurtosis(v, bias=False)}
    return tabla, params, preguntas, medidas, dispersion


def esperados_ej3(df, xs):
    x, y = df[X].to_numpy(), df[Y].to_numpy()
    b, a = np.polyfit(x, y, 1)
    r = stats.pearsonr(x, y).statistic
    nivel = "Perfecta" if abs(r) == 1 else "Excelente" if abs(r) >= .9 else "Aceptable" if abs(r) >= .8 else "Regular" if abs(r) >= .5 else "Mínima"
    reg = {"Pendiente (b)": b, "Intercepto (a)": a, "Ecuación": f"y = {round(a, 4)} {'+' if b >= 0 else '-'} {abs(round(b, 4))} x",
           "Tipo de relación": "Positiva" if b > 0 else "Negativa", "Promedio con x = 0": y[x == 0].mean()}
    cor = {"Coeficiente de determinación (R²)": r ** 2, "Confiabilidad del modelo (%)": r ** 2 * 100,
           "Coeficiente de correlación de Pearson (r)": r, "Nivel de correlación lineal": nivel}
    pred = {}
    for i, xv in enumerate(xs, start=1):
        m = x == xv
        pred.update({f"pred{i}": a + b * xv, f"obs{i}": y[m].mean(), f"n{i}": int(m.sum()), f"dentro{i}": "Sí" if x.min() <= xv <= x.max() else "No"})
    pred["error"] = np.sqrt(np.sum((y - (a + b * x)) ** 2) / (len(x) - 2))
    return reg, cor, pred


def comparar(filas, ejercicio, bloque, esperado, hoja_lo, hoja_guardada, celdas, nombres=None):
    for clave, esp in esperado.items():
        rc = celdas[clave]
        lo, guard = hoja_lo.cell(*rc).value, hoja_guardada.cell(*rc).value
        if isinstance(esp, str):
            dif, ok, ok_g = (0.0 if lo == esp else None), lo == esp, guard == esp
        else:
            dif = abs(float(lo) - float(esp))
            ok, ok_g = dif < TOL, abs(float(guard) - float(esp)) < TOL
        etiqueta = (nombres or {}).get(clave, clave)
        filas.append({"ejercicio": ejercicio, "bloque": bloque, "medida": etiqueta, "excel_libreoffice": lo, "python": esp if isinstance(esp, str) else float(esp),
                      "diferencia": dif, "ok": bool(ok), "valor_guardado_ok": bool(ok_g)})


def main():
    df = cargar_datos()
    filas = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        # Ejercicio 2: copia sin valores guardados (para recalcular) y copia con valores guardados (los que verá quien abra el archivo)
        pos, info = construir_ej2(tmp / "ej2_sin.xlsx", df, valores=False)
        construir_ej2(tmp / "ej2_con.xlsx", df, valores=True)
        lo2 = openpyxl.load_workbook(recalcular_con_libreoffice(tmp / "ej2_sin.xlsx", tmp), data_only=True)["Ej2 Tabla"]
        gu2 = openpyxl.load_workbook(tmp / "ej2_con.xlsx", data_only=True)["Ej2 Tabla"]
        tabla, params, preguntas, medidas, dispersion = esperados_ej2(df)
        # Tabla de frecuencia: 9 columnas por 10 intervalos, y los 6 parámetros de construcción
        celdas_t = {(nom, i): (pos["primera"] + 1 + i, 2 + j) for j, nom in enumerate(tabla) for i in range(len(tabla["fi"]))}
        esp_t = {(nom, i): tabla[nom][i] for nom in tabla for i in range(len(tabla["fi"]))}
        nombres_t = {(nom, i): f"Tabla, intervalo {i + 1}: {nom}" for nom in tabla for i in range(len(tabla["fi"]))}
        comparar(filas, "Ejercicio 2", "Tabla de frecuencia", esp_t, lo2, gu2, celdas_t, nombres_t)
        celdas_p = {nom: (3 + j, 2) for j, nom in enumerate(params)}
        comparar(filas, "Ejercicio 2", "Parámetros de la tabla", params, lo2, gu2, celdas_p, {k: f"Parámetro {k}" for k in params})
        comparar(filas, "Ejercicio 2", "Preguntas", preguntas, lo2, gu2, info["info_preguntas"]["celdas"])
        comparar(filas, "Ejercicio 2", "Tendencia central y posición", medidas, lo2, gu2, info["info_medidas"]["celdas"])
        comparar(filas, "Ejercicio 2", "Dispersión, asimetría y curtosis", dispersion, lo2, gu2, info["celdas"])
        # Ejercicio 3
        from ej3_predicciones import PREDICCIONES
        info3 = construir_ej3(tmp / "ej3_sin.xlsx", df, valores=False)
        construir_ej3(tmp / "ej3_con.xlsx", df, valores=True)
        lo3 = openpyxl.load_workbook(recalcular_con_libreoffice(tmp / "ej3_sin.xlsx", tmp), data_only=True)["Ej3 Regresión"]
        gu3 = openpyxl.load_workbook(tmp / "ej3_con.xlsx", data_only=True)["Ej3 Regresión"]
        reg, cor, pred = esperados_ej3(df, PREDICCIONES)
        comparar(filas, "Ejercicio 3", "Regresión lineal", reg, lo3, gu3, info3["celdas_regresion"])
        comparar(filas, "Ejercicio 3", "R², r y nivel", cor, lo3, gu3, info3["celdas_correlacion"])
        nombres_p = {f"{p}{i}": f"Predicción {i}: {t}" for i in range(1, 4)
                     for p, t in (("pred", "ŷ"), ("obs", "promedio observado"), ("n", "conductores"), ("dentro", "dentro del rango"))}
        nombres_p["error"] = "Error típico de la estimación"
        comparar(filas, "Ejercicio 3", "Predicciones", pred, lo3, gu3, info3["celdas"], nombres_p)
        # Hoja DATOS del libro contra el Anexo 1, celda por celda
        anexo = list(openpyxl.load_workbook(RUTA_ANEXO)["DATOS"].iter_rows(values_only=True))
        libro = list(openpyxl.load_workbook(tmp / "ej3_con.xlsx")["DATOS"].iter_rows(values_only=True))
        libro2 = list(openpyxl.load_workbook(tmp / "ej2_con.xlsx")["DATOS"].iter_rows(values_only=True))
    datos_ok = anexo == libro == libro2
    SALIDA.parent.mkdir(exist_ok=True)
    resumen = {"tolerancia": TOL, "filas": filas, "hoja_datos_identica_al_anexo": datos_ok, "registros_datos": len(anexo) - 1}
    SALIDA.write_text(json.dumps(resumen, ensure_ascii=False, indent=1, default=str), encoding="utf-8")

    n_ok = sum(f["ok"] for f in filas)
    n_g = sum(f["valor_guardado_ok"] for f in filas)
    difs = [f["diferencia"] for f in filas if f["diferencia"] is not None]
    print(f"Comparaciones: {len(filas)}; coinciden con numpy/scipy tras recalcular en LibreOffice: {n_ok}; valores guardados en el archivo que coinciden: {n_g}")
    print(f"Diferencia máxima absoluta (valores numéricos): {max(difs):.2e} (tolerancia {TOL:g})")
    print(f"Hoja DATOS idéntica al Anexo 1 en ambos libros: {datos_ok} ({len(anexo) - 1} registros)")
    por_bloque = {}
    for f in filas:
        por_bloque.setdefault((f["ejercicio"], f["bloque"]), []).append(f)
    for (e, b), fs in por_bloque.items():
        d = max((x["diferencia"] or 0) for x in fs)
        print(f"  {e} | {b}: {len(fs)} valores, todos OK: {all(x['ok'] for x in fs)}, dif. máx. {d:.1e}")
    fallos = [f for f in filas if not f["ok"] or not f["valor_guardado_ok"]]
    for f in fallos:
        print("  FALLO:", f)
    assert not fallos and datos_ok
    print(f"Tabla de verificación guardada en {SALIDA.relative_to(BASE)}")


if __name__ == "__main__":
    main()
